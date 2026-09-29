from pathlib import Path
import shutil
import sys

from openpyxl import load_workbook

OLD_VERSION = "A.1"
NEW_VERSION = "A.2"
EQUIPMENT_SHEET = "Equipment"
SYSTEM_SHEET = "_System"


def migrate(source: Path, target: Path) -> None:
    if source.resolve() == target.resolve():
        raise ValueError("Source and target must be different; original database must remain untouched.")

    shutil.copy2(source, target)
    wb = load_workbook(target)

    if SYSTEM_SHEET not in wb.sheetnames or EQUIPMENT_SHEET not in wb.sheetnames:
        raise ValueError("Database contract is missing _System or Equipment sheet.")

    ws_system = wb[SYSTEM_SHEET]
    version_found = False
    for row in ws_system.iter_rows():
        for idx, cell in enumerate(row):
            if cell.value == "schema_version" and idx + 1 < len(row) and row[idx + 1].value is not None:
                current = str(row[idx + 1].value)
                if current != OLD_VERSION:
                    raise ValueError(f"Expected schema {OLD_VERSION}, found {current}.")
                row[idx + 1].value = NEW_VERSION
                version_found = True
                break
        if version_found:
            break
    if not version_found:
        raise ValueError("schema_version entry not found in _System.")

    ws = wb[EQUIPMENT_SHEET]
    headers = [cell.value for cell in ws[1]]
    if headers == ["equipment_id", "unit_no", "category", "type", "owner_type", "owner_name", "status", "effective_from", "effective_to"]:
        wb.save(target)
        return

    expected_old = ["equipment_id", "category", "type", "owner_type", "owner_name", "status", "effective_from", "effective_to"]
    if headers != expected_old:
        raise ValueError(f"Unexpected Equipment headers: {headers}")

    ws.insert_cols(2, 1)
    ws.cell(1, 2).value = "unit_no"
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).value = None

    wb.save(target)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python tools/migrate_equipment_unit_no.py <source.xlsx> <target.xlsx>")
    migrate(Path(sys.argv[1]), Path(sys.argv[2]))
    print(f"Migrated database copy created: {sys.argv[2]}")
