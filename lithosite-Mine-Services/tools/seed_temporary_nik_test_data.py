"""One-time helper to seed explicitly temporary NIK demo records into the local A3 workbook.

Run from the Mine Services project root:
    python tools/seed_temporary_nik_test_data.py

The script creates a timestamped backup before changing the workbook. Remove the
TEST-NIK-* rows from the NIK sheet after validating the List NIK and Owner Name flow.
"""
from datetime import datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "Database" / "Mine-Services-Database-A3.xlsx"
EXPECTED_HEADERS = [
    "Effective From",
    "Effective To",
    "Id No (NIK)",
    "Origin",
    "Name Employment",
    "Information",
]
NAMES = [
    "Ari Waria",
    "Ari Nenen",
    "Ari Wakanda",
    "Ari Ucup Cup",
    "Aria Ozawa",
    "Ari Isap",
    "Ari Suntung",
    "Ari Itong",
]


def main():
    if not WORKBOOK_PATH.is_file():
        raise SystemExit(f"A3 workbook not found: {WORKBOOK_PATH}")

    workbook = load_workbook(WORKBOOK_PATH)
    try:
        if "NIK" not in workbook.sheetnames:
            raise SystemExit("Sheet NIK is missing. Start Desktop Host once to run the A3 schema migration.")
        sheet = workbook["NIK"]
        headers = [sheet.cell(row=1, column=column).value for column in range(1, 7)]
        if headers != EXPECTED_HEADERS:
            raise SystemExit(f"NIK headers do not match the contract: {headers!r}")

        existing = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = list(row) + [None] * max(0, 6 - len(row))
            nik = str(values[2] or "").strip()
            if nik:
                existing[nik.casefold()] = str(values[4] or "").strip()

        rows_to_add = []
        for index, name in enumerate(NAMES, start=1):
            nik = f"TEST-NIK-{index:03d}"
            current_name = existing.get(nik.casefold())
            if current_name:
                if current_name != name:
                    raise SystemExit(
                        f"{nik} already exists for {current_name!r}; refusing to overwrite it."
                    )
                continue
            rows_to_add.append([None, None, nik, "TEST", name, "TEMPORARY TEST DATA — remove after validation"])

        if not rows_to_add:
            print("All eight temporary NIK test records already exist; no changes made.")
            return

        archive = ROOT / "Database" / "Archive"
        archive.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup_path = archive / f"Mine-Services-Database-A3.pre-temporary-nik-test-{timestamp}.xlsx"
        shutil.copy2(WORKBOOK_PATH, backup_path)

        for row in rows_to_add:
            sheet.append(row)

        workbook.save(WORKBOOK_PATH)
        print(f"Added {len(rows_to_add)} temporary NIK test records.")
        print(f"Workbook: {WORKBOOK_PATH}")
        print(f"Backup:   {backup_path}")
        print("IDs used: TEST-NIK-001 through TEST-NIK-008")
        print("After validation, delete rows with TEST-NIK-* from the NIK sheet.")
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
