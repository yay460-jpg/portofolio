from copy import deepcopy
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook

from . import schema as DEFAULT_SCHEMA


class PersistenceStore:
    def __init__(self, path=None, schema_module=None):
        self.schema = schema_module or DEFAULT_SCHEMA
        self.path = Path(path) if path else None
        self._data = {e: {} for e in self.schema.DOMAIN_ENTITIES}
        self._audit = []
        self.controlled_lists = {
            "equipment_category": {"Heavy Equipment", "Light Vehicle", "Support Equipment"},
            "equipment_type": {"Dump Truck", "Excavator", "Dozer", "Grader", "Water Truck", "Loader", "Light Vehicle", "Other"},
            "owner_type": {"Owner", "Contractor"},
            "equipment_status": {"Active", "Inactive", "Retired"},
            "service_domain": {"Road & Hauling", "Drainage & Dewatering", "Land Clearing", "Disposal & Stockpile", "Mining", "Reclamation", "Other"},
            "work_front_status": {"Active", "Inactive", "Closed"},
            "capacity_status": {"Active", "Inactive"},
            "checker_shift": {"Day", "Night"},
            "checker_material": {"Ore", "OB", "Quarry"},
            "transaction_status": {"DRAFT", "VALIDATED", "REJECTED", "VOIDED"},
            "unit": {"hour", "km", "m", "m2", "m3", "ton", "unit", "cycle"},
            "maintenance_event_type": {"Preventive", "Corrective", "Inspection", "Breakdown"},
            "issue_severity": {"Low", "Medium", "High", "Critical"},
            "issue_status": {"Open", "In Progress", "Closed", "Void"},
            "hse_event_type": {"Inspection", "Incident", "Near Miss", "Environmental", "Corrective Action"},
            "maintenance_status": {"Open", "In Progress", "Completed", "Cancelled"},
            "plan_status": {"Draft", "Approved", "In Progress", "Completed", "Cancelled"},
            "hse_severity": {"Low", "Medium", "High", "Critical"},
            "hse_status": {"Open", "In Progress", "Closed", "Void"},
        }
        if self.path and self.path.exists():
            self.load()

    def snapshot(self):
        return {"data": deepcopy(self._data), "audit": deepcopy(self._audit)}

    def replace(self, snapshot):
        self._data = deepcopy(snapshot["data"])
        self._audit = deepcopy(snapshot.get("audit", []))

    def exists(self, entity, pk):
        return pk in self._data.get(entity, {})

    def get(self, entity, pk):
        return deepcopy(self._data.get(entity, {}).get(pk))

    def all(self, entity):
        return deepcopy(list(self._data.get(entity, {}).values()))

    def insert(self, entity, pk, row):
        self._data.setdefault(entity, {})[pk] = deepcopy(dict(row))

    def update(self, entity, pk, row):
        if not self.exists(entity, pk):
            raise KeyError(pk)
        self._data[entity][pk] = deepcopy(dict(row))

    def delete(self, entity, pk):
        if not self.exists(entity, pk):
            raise KeyError(pk)
        del self._data[entity][pk]

    def add_audit(self, event):
        self._audit.append(deepcopy(event))

    def audit(self):
        return deepcopy(self._audit)

    @staticmethod
    def _system_version(workbook):
        if "_System" not in workbook.sheetnames:
            raise ValueError("SHEET_MISSING:_System")
        for values in workbook["_System"].iter_rows(values_only=True):
            vals = list(values)
            for i, value in enumerate(vals):
                if value == "schema_version" and i + 1 < len(vals) and vals[i + 1] is not None:
                    return str(vals[i + 1])
        raise ValueError("SCHEMA_VERSION:MISSING")

    def _migrate_a3_checker_support_headers(self, workbook):
        """Upgrade the existing A3 workbook headers for the V38 Operations/Checker support layer.

        This is intentionally narrow: only the known pre-support A3 headers are migrated.
        Existing row values are preserved; new support fields start blank.
        """
        if getattr(self.schema, "SCHEMA_VERSION", None) != "A.3":
            return False

        changed = False

        operations = workbook["Operations"]
        current_operations = [cell.value for cell in operations[1]]
        legacy_operations = list(self.schema.HEADERS["Operations"][:-4])
        if current_operations == legacy_operations:
            for header in self.schema.HEADERS["Operations"][-4:]:
                operations.cell(row=1, column=operations.max_column + 1, value=header)
            changed = True

        checker = workbook["Checker"]
        current_checker = [cell.value for cell in checker[1]]
        legacy_checker = [h for h in self.schema.HEADERS["Checker"] if h != "operation_id"]
        if current_checker == legacy_checker:
            checker.insert_cols(2, 1)
            checker.cell(row=1, column=2, value="operation_id")
            changed = True

        return changed

    def _validate_workbook_contract(self, workbook):
        required = set(self.schema.DOMAIN_ENTITIES) | {"_System", "_Lists", "AuditLog"}
        missing = required - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"SHEET_MISSING:{sorted(missing)}")

        if self._system_version(workbook) != self.schema.SCHEMA_VERSION:
            raise ValueError(f"SCHEMA_VERSION:{self._system_version(workbook)}")

        for entity in self.schema.DOMAIN_ENTITIES:
            rows = list(workbook[entity].values)
            if not rows or list(rows[0]) != self.schema.HEADERS[entity]:
                raise ValueError(f"HEADER_MISMATCH:{entity}")

        audit_rows = list(workbook["AuditLog"].values)
        if not audit_rows or list(audit_rows[0]) != self.schema.HEADERS["AuditLog"]:
            raise ValueError("HEADER_MISMATCH:AuditLog")

    def load(self):
        wb = load_workbook(self.path, data_only=True)
        migrated = self._migrate_a3_checker_support_headers(wb)
        if migrated:
            wb.save(self.path)
            wb = load_workbook(self.path, data_only=True)
        self._validate_workbook_contract(wb)

        for entity in self.schema.DOMAIN_ENTITIES:
            rows = list(wb[entity].values)
            self._data[entity] = {}
            pk = self.schema.PKS[entity]
            for values in rows[1:]:
                if all(value is None for value in values):
                    continue
                row = dict(zip(self.schema.HEADERS[entity], values))
                self._data[entity][row[pk]] = row

        values = list(wb["_Lists"].values)
        if values:
            workbook_lists = {
                header: {
                    row[i]
                    for row in values[1:]
                    if i < len(row) and row[i] not in (None, "")
                }
                for i, header in enumerate(values[0])
                if header not in (None, "")
            }
            workbook_lists.setdefault("service_domain", set()).add("Mining")
            workbook_lists.setdefault("capacity_status", set()).update({"Active", "Inactive"})
            workbook_lists.setdefault("unit", set()).add("cycle")
            workbook_lists.setdefault("checker_shift", set()).update({"Day", "Night"})
            workbook_lists.setdefault("checker_material", set()).update({"Ore", "OB", "Quarry"})
            self.controlled_lists = workbook_lists

        rows = list(wb["AuditLog"].values)
        self._audit = [
            dict(zip(self.schema.HEADERS["AuditLog"], values))
            for values in rows[1:]
            if not all(value is None for value in values)
        ]

    def save(self):
        if not self.path:
            return

        wb = load_workbook(self.path)
        self._validate_workbook_contract(wb)

        for entity in self.schema.DOMAIN_ENTITIES:
            ws = wb[entity]
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            for row in self.all(entity):
                ws.append([row.get(header) for header in self.schema.HEADERS[entity]])

        ws = wb["AuditLog"]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for row in self.audit():
            ws.append([row.get(header) for header in self.schema.HEADERS["AuditLog"]])

        with NamedTemporaryFile(
            prefix=self.path.stem + ".",
            suffix=self.path.suffix,
            dir=self.path.parent,
            delete=False,
        ) as file:
            temp_path = Path(file.name)

        try:
            wb.save(temp_path)
            temp_path.replace(self.path)
        finally:
            if temp_path.exists():
                temp_path.unlink()
