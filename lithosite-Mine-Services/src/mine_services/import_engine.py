from openpyxl import load_workbook

from .schema import DOMAIN_ENTITIES, HEADERS, PKS, SCHEMA_VERSION
from .validation import ValidationEngine
from .persistence import PersistenceStore
from .audit import AuditRepository

IMPORT_ERROR_CODES = {
    "FILE_UNREADABLE": "E001",
    "SCHEMA_VERSION": "E002",
    "SHEET_MISSING": "E003",
    "SHEET_UNEXPECTED": "E004",
    "HEADER_MISMATCH": "E005",
    "TYPE_INVALID": "E006",
    "REQUIRED_MISSING": "E007",
    "PK_DUPLICATE": "E008",
    "FK_NOT_FOUND": "E009",
    "ENUM_INVALID": "E010",
    "RANGE_INVALID": "E011",
    "CONDITIONAL_INVALID": "E012",
    "PERIOD_INVALID": "E013",
    "ATOMIC_ABORT": "E014",
}

class ImportCoordinator:
    def __init__(self, store, validator=None, transaction=None, audit_repository=None):
        from .transaction import TransactionManager
        self.store = store
        self.validator = validator or ValidationEngine()
        self.transaction = transaction or TransactionManager(store)
        self.audit_repository = audit_repository or AuditRepository(store)

    @staticmethod
    def _error(code, message, field=None):
        return {"code": code, "field": field, "message": message}

    @staticmethod
    def _map_validation_error(error):
        mapping = {
            "VAL-E002": IMPORT_ERROR_CODES["TYPE_INVALID"],
            "VAL-E003": IMPORT_ERROR_CODES["REQUIRED_MISSING"],
            "VAL-E004": IMPORT_ERROR_CODES["PK_DUPLICATE"],
            "VAL-E005": IMPORT_ERROR_CODES["FK_NOT_FOUND"],
            "VAL-E006": IMPORT_ERROR_CODES["ENUM_INVALID"],
            "VAL-E007": IMPORT_ERROR_CODES["RANGE_INVALID"],
            "VAL-E008": IMPORT_ERROR_CODES["CONDITIONAL_INVALID"],
            "VAL-E009": IMPORT_ERROR_CODES["PERIOD_INVALID"],
            "VAL-E010": IMPORT_ERROR_CODES["REQUIRED_MISSING"],
            "VAL-E011": IMPORT_ERROR_CODES["HEADER_MISMATCH"],
        }
        return {
            "code": mapping.get(error.code, IMPORT_ERROR_CODES["ATOMIC_ABORT"]),
            "field": error.field,
            "message": error.message,
        }

    @staticmethod
    def _schema_version_from_system(ws):
        values = list(ws.values)
        for row in values:
            for index, value in enumerate(row):
                if value == "schema_version" and index + 1 < len(row):
                    candidate = row[index + 1]
                    if candidate is not None:
                        return str(candidate)
        for row in values:
            if len(row) == 1 and row[0] is not None:
                text = str(row[0])
                if text == SCHEMA_VERSION:
                    return text
        return None

    def read_xlsx(self, path):
        try:
            wb = load_workbook(path, data_only=True)
        except Exception as exc:
            return None, [self._error(IMPORT_ERROR_CODES["FILE_UNREADABLE"], f"Workbook could not be opened: {exc}")]

        required = set(DOMAIN_ENTITIES) | {"_System", "_Lists"}
        missing = required - set(wb.sheetnames)
        if missing:
            return None, [self._error(IMPORT_ERROR_CODES["SHEET_MISSING"], f"Missing sheets: {sorted(missing)}")]

        allowed = required | {"_Baseline", "AuditLog"}
        unexpected = set(wb.sheetnames) - allowed
        if unexpected:
            return None, [self._error(IMPORT_ERROR_CODES["SHEET_UNEXPECTED"], f"Unexpected sheets: {sorted(unexpected)}")]

        version = self._schema_version_from_system(wb["_System"])
        if version != SCHEMA_VERSION:
            return None, [self._error(IMPORT_ERROR_CODES["SCHEMA_VERSION"], f"Expected schema {SCHEMA_VERSION}, got {version}")]

        rows = {}
        for entity in DOMAIN_ENTITIES:
            values = list(wb[entity].values)
            if not values:
                return None, [self._error(IMPORT_ERROR_CODES["HEADER_MISMATCH"], "Sheet is empty", entity)]
            headers = list(values[0])
            if headers != HEADERS[entity]:
                return None, [self._error(IMPORT_ERROR_CODES["HEADER_MISMATCH"], "Header mismatch", entity)]
            rows[entity] = [
                dict(zip(headers, values_row))
                for values_row in values[1:]
                if not all(value is None for value in values_row)
            ]
        return rows, []

    def import_xlsx(self, path):
        rows, errors = self.read_xlsx(path)
        if errors:
            return {"status": "REJECTED", "lifecycle": "REJECTED", "rows_committed": 0, "errors": errors}
        return self.import_dataset(rows)

    def import_dataset(self, rows):
        staged = PersistenceStore()
        staged.replace(self.store.snapshot())

        duplicate_errors = []
        seen = {}
        for entity, items in rows.items():
            if entity not in DOMAIN_ENTITIES:
                duplicate_errors.append(self._error(IMPORT_ERROR_CODES["HEADER_MISMATCH"], f"Unknown entity: {entity}", entity))
                continue
            pk = PKS[entity]
            seen.setdefault(entity, set())
            for row in items:
                key = row.get(pk)
                if key in seen[entity]:
                    duplicate_errors.append(self._error(IMPORT_ERROR_CODES["PK_DUPLICATE"], "Primary key duplicate in dataset", pk))
                seen[entity].add(key)
                if key not in (None, "") and staged.exists(entity, key):
                    duplicate_errors.append(self._error(IMPORT_ERROR_CODES["PK_DUPLICATE"], "Primary key already exists", pk))

        if duplicate_errors:
            return {
                "status": "REJECTED",
                "lifecycle": "REJECTED",
                "rows_committed": 0,
                "errors": duplicate_errors,
            }

        try:
            for entity, items in rows.items():
                for row in items:
                    staged.insert(entity, row[PKS[entity]], row)

            errors = self.validator.validate_dataset(rows, staged, "IMPORT")
            if errors:
                return {
                    "status": "REJECTED",
                    "lifecycle": "REJECTED",
                    "rows_committed": 0,
                    "errors": [self._map_validation_error(error) for error in errors],
                }

            self.transaction.begin()
            try:
                for entity, items in rows.items():
                    for row in items:
                        self.store.insert(entity, row[PKS[entity]], row)

                self.audit_repository.append(
                    entity="_System",
                    entity_id="IMPORT",
                    action="IMPORT",
                    request_id="IMPORT",
                    old_value=None,
                    new_value=str({entity: len(items) for entity, items in rows.items()}),
                    source="ImportCoordinator",
                )
                self.transaction.commit()
                return {
                    "status": "COMMITTED",
                    "lifecycle": "AUDITED",
                    "rows_committed": sum(len(items) for items in rows.values()),
                    "errors": [],
                }
            except Exception:
                self.transaction.rollback()
                return {
                    "status": "REJECTED",
                    "lifecycle": "REJECTED",
                    "rows_committed": 0,
                    "errors": [self._error(IMPORT_ERROR_CODES["ATOMIC_ABORT"], "Atomic import aborted")],
                }
        except Exception:
            return {
                "status": "REJECTED",
                "lifecycle": "REJECTED",
                "rows_committed": 0,
                "errors": [self._error(IMPORT_ERROR_CODES["ATOMIC_ABORT"], "Atomic import aborted")],
            }
