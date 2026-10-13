"""One-time A.3 repair for GlobalCapacity, Checker, and hauling field headers.

The migration accepts an A.3 workbook with the known canonical header prefixes
for WorkFront, Operations, GlobalCapacity, and Checker. It refuses to upgrade
an unknown schema or overwrite data in unnamed columns.
"""
from pathlib import Path
import sys
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook

from . import schema

DEFAULT_DB = Path(__file__).resolve().parents[2] / "Database" / "Mine-Services-Database-A3.xlsx"
REPAIRABLE_SHEETS = ("GlobalCapacity", "Checker", "WorkFront", "Operations")
CREATABLE_SHEETS = {"GlobalCapacity", "Checker"}


def _read_headers(worksheet):
    headers = [cell.value for cell in worksheet[1]]
    while headers and headers[-1] in (None, ""):
        headers.pop()
    return headers


def _read_schema_version(workbook):
    worksheet = workbook["_System"]
    for row in worksheet.iter_rows(values_only=True):
        values = list(row)
        for index, value in enumerate(values):
            if value == "schema_version":
                if index + 1 >= len(values) or values[index + 1] in (None, ""):
                    raise ValueError("SCHEMA_VERSION:VALUE_MISSING")
                return str(values[index + 1])
    raise ValueError("SCHEMA_VERSION:MISSING")


def _ensure_header_prefix(worksheet, expected):
    current = _read_headers(worksheet)
    if len(current) > len(expected) or current != list(expected[:len(current)]):
        raise ValueError(
            f"HEADER_MISMATCH:{worksheet.title}; expected a canonical header prefix, got {current!r}"
        )

    # A field with no header cannot safely be assigned to a canonical column.
    known_width = len(current)
    for row_number in range(2, worksheet.max_row + 1):
        for column in range(known_width + 1, worksheet.max_column + 1):
            value = worksheet.cell(row=row_number, column=column).value
            if value not in (None, ""):
                raise ValueError(
                    f"UNMAPPED_COLUMN_DATA:{worksheet.title}:row={row_number}:column={column}"
                )

    for header in expected[len(current):]:
        current.append(header)
        worksheet.cell(row=1, column=len(current), value=header)
    return current


def _verify_target(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        missing = set(schema.SHEETS) - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"MIGRATION_VERIFY_FAILED:SHEETS:{sorted(missing)}")
        if _read_schema_version(workbook) != schema.SCHEMA_VERSION:
            raise ValueError("MIGRATION_VERIFY_FAILED:SCHEMA_VERSION")

        for sheet in REPAIRABLE_SHEETS:
            actual = _read_headers(workbook[sheet])
            expected = list(schema.HEADERS[sheet])
            if actual != expected:
                raise ValueError(f"MIGRATION_VERIFY_FAILED:HEADER_MISMATCH:{sheet}")
    finally:
        workbook.close()


def migrate(path: Path):
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)

    workbook = load_workbook(path)
    temp_path = None
    try:
        missing = set(schema.SHEETS) - set(workbook.sheetnames)
        unrepairable_missing = missing - CREATABLE_SHEETS
        if unrepairable_missing:
            raise ValueError(f"Cannot migrate: required sheet is missing: {sorted(unrepairable_missing)}")

        # This helper repairs known incomplete A.3 headers; A.2 must use the
        # dedicated A.2-to-A.3 migration rather than being relabelled here.
        current_version = _read_schema_version(workbook)
        if current_version != schema.SCHEMA_VERSION:
            raise ValueError(
                f"SCHEMA_VERSION:{current_version} (expected {schema.SCHEMA_VERSION})"
            )

        for sheet in ("GlobalCapacity", "Checker"):
            if sheet not in workbook.sheetnames:
                worksheet = workbook.create_sheet(sheet, index=schema.SHEETS.index(sheet))
                worksheet.append(schema.HEADERS[sheet])
            else:
                _ensure_header_prefix(workbook[sheet], schema.HEADERS[sheet])

        for sheet in ("WorkFront", "Operations"):
            _ensure_header_prefix(workbook[sheet], schema.HEADERS[sheet])

        # Write to a sibling workbook and verify it before replacing the live
        # database, so validation/save/replace failures preserve the source.
        with NamedTemporaryFile(
            prefix=path.stem + ".",
            suffix=path.suffix,
            dir=path.parent,
            delete=False,
        ) as handle:
            temp_path = Path(handle.name)

        workbook.save(temp_path)
        workbook.close()
        _verify_target(temp_path)
        temp_path.replace(path)
        return path
    finally:
        workbook.close()
        if temp_path is not None and temp_path.exists():
            temp_path.unlink()


if __name__ == "__main__":
    target = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_DB
    print(migrate(target))
