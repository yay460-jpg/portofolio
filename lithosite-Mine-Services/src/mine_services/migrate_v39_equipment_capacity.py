"""One-time V39 migration: move operational capacity ownership from WorkFront to Equipment.

The migration is intentionally conservative:
- Existing Equipment.capacity_profile_id values are preserved.
- Blank Equipment profiles are inferred only from Operations history.
- A profile is assigned only when exactly one valid Global Capacity profile can
  be resolved for that equipment.
- Ambiguous or unresolved equipment remains blank for manual assignment.

Run:
    python -m src.mine_services.migrate_v39_equipment_capacity [path-to-xlsx]
"""
from pathlib import Path
import sys

from openpyxl import load_workbook

from . import schema


DEFAULT_DB = Path(__file__).resolve().parents[2] / "Database" / "Mine-Services-Database-A3.xlsx"


def migrate(path: Path):
    wb = load_workbook(path)
    for sheet in ("Equipment", "Operations", "WorkFront", "GlobalCapacity"):
        if sheet not in wb.sheetnames:
            raise ValueError(f"Cannot migrate: required sheet is missing: {sheet}")

    equipment_ws = wb["Equipment"]
    operations_ws = wb["Operations"]
    workfront_ws = wb["WorkFront"]
    capacity_ws = wb["GlobalCapacity"]

    def read_headers(worksheet):
        headers = [cell.value for cell in worksheet[1]]
        while headers and headers[-1] in (None, ""):
            headers.pop()
        return headers

    equipment_headers = read_headers(equipment_ws)
    current_equipment_headers = list(schema.HEADERS["Equipment"])
    legacy_equipment_headers = current_equipment_headers[:-1]

    # V38/A3 workbooks may not yet contain the V39 Equipment profile field.
    # Accept only the exact known legacy header, and reject any populated
    # unnamed columns rather than risking overwriting user data.
    if equipment_headers == legacy_equipment_headers:
        legacy_width = len(legacy_equipment_headers)
        if equipment_ws.max_column > legacy_width:
            for row_number in range(2, equipment_ws.max_row + 1):
                if any(
                    equipment_ws.cell(row=row_number, column=column).value not in (None, "")
                    for column in range(legacy_width + 1, equipment_ws.max_column + 1)
                ):
                    raise ValueError(
                        "HEADER_MISMATCH:Equipment has data beyond the recognized legacy header"
                    )
        equipment_ws.cell(row=1, column=legacy_width + 1, value="capacity_profile_id")
        equipment_headers = read_headers(equipment_ws)
    elif equipment_headers != current_equipment_headers:
        raise ValueError(
            "HEADER_MISMATCH:Equipment; expected current V39 header or exact pre-V39 A3 header; "
            f"found {equipment_headers!r}"
        )

    operation_headers = read_headers(operations_ws)
    workfront_headers = read_headers(workfront_ws)
    capacity_headers = read_headers(capacity_ws)

    required = {
        "Operations": (operation_headers, {"equipment_id", "capacity_profile_id", "work_front_id"}),
        "WorkFront": (workfront_headers, {"work_front_id", "capacity_profile_id"}),
        "GlobalCapacity": (capacity_headers, {"capacity_profile_id"}),
    }
    for sheet, (headers, fields) in required.items():
        missing = sorted(fields - set(headers))
        if missing:
            raise ValueError(f"HEADER_MISMATCH:{sheet}; missing required fields: {', '.join(missing)}")

    equipment_index = {name: index for index, name in enumerate(equipment_headers)}
    operation_index = {name: index for index, name in enumerate(operation_headers)}
    workfront_index = {name: index for index, name in enumerate(workfront_headers)}
    capacity_index = {name: index for index, name in enumerate(capacity_headers)}

    valid_profiles = {
        row[capacity_index["capacity_profile_id"]]
        for row in capacity_ws.iter_rows(min_row=2, values_only=True)
        if row[capacity_index["capacity_profile_id"]] not in (None, "")
    }

    workfront_profiles = {
        row[workfront_index["work_front_id"]]: row[workfront_index["capacity_profile_id"]]
        for row in workfront_ws.iter_rows(min_row=2, values_only=True)
        if row[workfront_index["work_front_id"]] not in (None, "")
    }

    operation_rows = list(operations_ws.iter_rows(min_row=2, values_only=True))
    resolved = []
    ambiguous = []
    unresolved = []
    preserved = []

    for row_number, row in enumerate(equipment_ws.iter_rows(min_row=2), start=2):
        equipment_id = row[equipment_index["equipment_id"]].value
        current_profile = row[equipment_index["capacity_profile_id"]].value
        if current_profile not in (None, ""):
            preserved.append((equipment_id, current_profile))
            continue

        candidates = set()
        for operation in operation_rows:
            if operation[operation_index["equipment_id"]] != equipment_id:
                continue

            profile = operation[operation_index["capacity_profile_id"]]
            if profile in valid_profiles:
                candidates.add(profile)

            if profile in (None, ""):
                legacy_workfront = workfront_profiles.get(operation[operation_index["work_front_id"]])
                if legacy_workfront in valid_profiles:
                    candidates.add(legacy_workfront)

        if len(candidates) == 1:
            profile = next(iter(candidates))
            row[equipment_index["capacity_profile_id"]].value = profile
            resolved.append((equipment_id, profile))
        elif len(candidates) > 1:
            ambiguous.append((equipment_id, sorted(candidates)))
        else:
            unresolved.append(equipment_id)

    wb.save(path)
    return resolved, ambiguous, unresolved, preserved


if __name__ == "__main__":
    target = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_DB
    resolved, ambiguous, unresolved, preserved = migrate(target)
    print(f"Resolved: {len(resolved)}")
    for equipment_id, profile in resolved:
        print(f"  {equipment_id} -> {profile}")
    print(f"Ambiguous: {len(ambiguous)}")
    for equipment_id, profiles in ambiguous:
        print(f"  {equipment_id} -> {', '.join(profiles)}")
    print(f"Already assigned/preserved: {len(preserved)}")
    for equipment_id, profile in preserved:
        print(f"  {equipment_id} -> {profile}")
    print(f"Unresolved: {len(unresolved)}")
    for equipment_id in unresolved:
        print(f"  {equipment_id} -> no unique valid capacity profile; manual review required")
    print(target)
