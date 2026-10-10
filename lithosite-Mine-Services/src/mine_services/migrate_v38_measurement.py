"""One-time migration: Unit -> Measurement.

Renames only schema field headers in the active A3 workbook.
Equipment.unit_no is intentionally untouched.

Run:
    python -m src.mine_services.migrate_measurement
or:
    python -m src.mine_services.migrate_measurement path/to/Mine-Services-Database-A3.xlsx
"""
from pathlib import Path
import sys

from openpyxl import load_workbook


DEFAULT_DB = (
    Path(__file__).resolve().parents[2]
    / "Database"
    / "Mine-Services-Database-A3.xlsx"
)

RENAMES = {
    "Operations": {"unit": "measurement", "capacity_unit": "capacity_measurement"},
    "Plans": {"unit": "measurement"},
    "GlobalCapacity": {"unit": "measurement"},
    "_Lists": {"unit": "measurement"},
}


def rename_headers(ws, mapping):
    headers = [cell.value for cell in ws[1]]
    for old, new in mapping.items():
        if old not in headers:
            if new in headers:
                continue
            raise ValueError(f"{ws.title}: missing expected header {old!r}")
        if new in headers:
            raise ValueError(f"{ws.title}: both {old!r} and {new!r} already exist")
        headers[headers.index(old)] = new
        ws.cell(row=1, column=headers.index(new) + 1, value=new)


def migrate(path):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    wb = load_workbook(path)
    changed = False
    for sheet, mapping in RENAMES.items():
        ws = wb[sheet]
        before = [cell.value for cell in ws[1]]
        rename_headers(ws, mapping)
        changed |= before != [cell.value for cell in ws[1]]

    if changed:
        wb.save(path)

    # Verify the resulting workbook against the canonical field names.
    check = load_workbook(path, read_only=True, data_only=False)
    expected = {
        "Operations": ["measurement", "capacity_measurement"],
        "Plans": ["measurement"],
        "GlobalCapacity": ["measurement"],
        "_Lists": ["measurement"],
    }
    for sheet, names in expected.items():
        headers = [cell.value for cell in check[sheet][1]]
        for name in names:
            if name not in headers:
                raise RuntimeError(f"Migration verification failed: {sheet}.{name}")
    print(f"Measurement migration PASS: {path}")


if __name__ == "__main__":
    migrate(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_DB)
