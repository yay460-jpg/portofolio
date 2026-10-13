import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))

from mine_services import (
    ApplicationService,
    ImportCoordinator,
    PersistenceStore,
    RuntimeInterface,
    SnapshotManager,
)


def valid_equipment(equipment_id="EQ-1"):
    return {
        "equipment_id": equipment_id,
        "unit_no": "DT-001",
        "category": "Heavy Equipment",
        "type": "Dump Truck",
        "owner_type": "Owner",
        "status": "Active",
    }


def test_validation_and_crud():
    app = ApplicationService(PersistenceStore())
    assert app.create("Equipment", valid_equipment(), "r1")["status"] == "COMMITTED"
    assert app.update("Equipment", "EQ-1", {"status": "Inactive"}, "r2")["status"] == "COMMITTED"
    assert app.delete("Equipment", "EQ-1", "r3")["status"] == "COMMITTED"
    assert len(app.store.audit()) == 3


def test_runtime_reads_audit_log_through_audit_repository():
    app = ApplicationService(PersistenceStore())
    runtime = RuntimeInterface(application=app)

    assert app.create("Equipment", valid_equipment("EQ-AUDIT"), "audit-create")["status"] == "COMMITTED"

    events = runtime.read("AuditLog")
    assert len(events) == 1
    assert events[0]["entity"] == "Equipment"
    assert events[0]["entity_id"] == "EQ-AUDIT"
    assert events[0]["action"] == "CREATE"
    assert events[0]["source"] == "audit-create"

    event = runtime.read("AuditLog", events[0]["audit_id"])
    assert event["audit_id"] == events[0]["audit_id"]


def test_fk_and_enum_rejected():
    app = ApplicationService()
    fk_result = app.create(
        "Maintenance",
        {
            "maintenance_id": "M1",
            "equipment_id": "NOPE",
            "event_type": "Preventive",
            "status": "Open",
        },
        "r1",
    )
    assert fk_result["status"] == "REJECTED"

    enum_result = app.create(
        "Equipment",
        {
            "equipment_id": "E1",
            "unit_no": "DT-001",
            "category": "Invalid Category",
            "type": "Dump Truck",
            "owner_type": "Owner",
            "status": "Active",
        },
        "r2",
    )
    assert enum_result["status"] == "REJECTED"


def test_import_atomic_duplicate():
    store = PersistenceStore()
    result = ImportCoordinator(store).import_dataset(
        {
            "Equipment": [
                valid_equipment("E"),
                valid_equipment("E"),
            ]
        }
    )
    assert result["status"] == "REJECTED"
    assert store.all("Equipment") == []


def test_snapshot_tamper():
    store = PersistenceStore()
    app = ApplicationService(store)
    assert app.create("Equipment", valid_equipment("E"), "r")["status"] == "COMMITTED"
    snap = SnapshotManager(store).capture()
    snap["checksum"] = "tampered"
    assert SnapshotManager(store).restore(snap)["status"] == "REJECTED"


def test_runtime_exposes_authoritative_controlled_lists():
    runtime = RuntimeInterface(application=ApplicationService(PersistenceStore()))
    lists = runtime.read("_Lists")
    assert lists["equipment_category"] == ["Heavy Equipment", "Light Vehicle", "Support Equipment"]
    assert lists["equipment_type"] == ["Dozer", "Dump Truck", "Excavator", "Grader", "Light Vehicle", "Loader", "Other", "Water Truck"]
    assert lists["owner_type"] == ["Contractor", "Owner"]
    assert lists["equipment_status"] == ["Active", "Inactive", "Retired"]


def valid_workfront(work_front_id="WF-1"):
    return {
        "work_front_id": work_front_id,
        "domain": "Road & Hauling",
        "location": "Pit North",
        "responsible": "Operations Team",
        "status": "Active",
    }


def test_workfront_crud_and_audit():
    app = ApplicationService(PersistenceStore())
    assert app.create("WorkFront", valid_workfront(), "wf-create")["status"] == "COMMITTED"
    assert app.update("WorkFront", "WF-1", {"location": "Pit South"}, "wf-update")["status"] == "COMMITTED"
    assert app.store.get("WorkFront", "WF-1")["location"] == "Pit South"
    assert app.delete("WorkFront", "WF-1", "wf-delete")["status"] == "COMMITTED"
    assert [event["action"] for event in app.store.audit()] == ["CREATE", "UPDATE", "DELETE"]


def test_workfront_delete_is_rejected_when_referenced():
    app = ApplicationService(PersistenceStore())
    assert app.create("WorkFront", valid_workfront("WF-REF"), "wf-ref")["status"] == "COMMITTED"
    operation = {
        "transaction_id": "TR-WF-REF",
        "transaction_date": "2026-09-29",
        "transaction_time": "08:00",
        "domain": "Road & Hauling",
        "work_front_id": "WF-REF",
        "activity": "Functional Test",
        "quantity": 10,
        "measurement": "ton",
        "actual_hours": 1,
        "target_hours": 2,
        "status": "VALIDATED",
        "source": "Stage10",
    }
    assert app.create("Operations", operation, "wf-operation")["status"] == "COMMITTED"
    result = app.delete("WorkFront", "WF-REF", "wf-delete-ref")
    assert result["status"] == "REJECTED"
    assert app.store.exists("WorkFront", "WF-REF")


def test_workfront_controlled_vocab_and_xlsx_roundtrip(tmp_path):
    from openpyxl import Workbook
    from mine_services.schema import DOMAIN_ENTITIES, HEADERS, SCHEMA_VERSION

    path = tmp_path / "workfront-stage10.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    baseline = wb.create_sheet("_Baseline")
    baseline.append(["baseline_status", "LOCKED"])
    system = wb.create_sheet("_System")
    system.append(["schema_version", SCHEMA_VERSION])
    lists = wb.create_sheet("_Lists")
    defaults = PersistenceStore().controlled_lists
    headers = list(defaults)
    lists.append(headers)
    for i in range(max(len(values) for values in defaults.values())):
        lists.append([
            sorted(defaults[h])[i] if i < len(defaults[h]) else None
            for h in headers
        ])
    for entity in DOMAIN_ENTITIES:
        ws = wb.create_sheet(entity)
        ws.append(HEADERS[entity])
    audit = wb.create_sheet("AuditLog")
    audit.append(HEADERS["AuditLog"])
    wb.save(path)

    store = PersistenceStore(path)
    app = ApplicationService(store)
    invalid = valid_workfront("WF-BAD")
    invalid["domain"] = "Not A Domain"
    assert app.create("WorkFront", invalid, "wf-invalid")["status"] == "REJECTED"

    assert app.create("WorkFront", valid_workfront("WF-XLSX"), "wf-xlsx")["status"] == "COMMITTED"
    reloaded = PersistenceStore(path)
    assert reloaded.get("WorkFront", "WF-XLSX")["domain"] == "Road & Hauling"
    assert "Road & Hauling" in reloaded.controlled_lists["service_domain"]

def test_runtime_dataset_save_overwrites_same_file_and_save_as_creates_new(tmp_path):
    app = ApplicationService(PersistenceStore())
    runtime = RuntimeInterface(application=app)
    runtime._datasets.directory = tmp_path

    assert app.create("Equipment", valid_equipment("E-DATASET"), "dataset-create")["status"] == "COMMITTED"

    first = runtime.save_dataset()
    second = runtime.save_dataset()
    assert first["filename"] == second["filename"]
    assert first["dataset_id"] == second["dataset_id"]
    assert len(list(tmp_path.glob("*.json"))) == 1

    saved_as = runtime.save_dataset_as("October Review")
    assert saved_as["dataset_name"] == "October Review"
    assert saved_as["dataset_id"] != first["dataset_id"]
    assert len(list(tmp_path.glob("*.json"))) == 2


def test_runtime_dataset_save_rolls_back_file_and_audit_when_store_save_fails(tmp_path):
    app = ApplicationService(PersistenceStore())
    runtime = RuntimeInterface(application=app)
    runtime._datasets.directory = tmp_path

    assert app.create("Equipment", valid_equipment("E-SAVE-ROLLBACK"), "save-rollback-create")["status"] == "COMMITTED"
    saved = runtime.save_dataset("Rollback State")
    target = tmp_path / saved["filename"]
    prior_file = target.read_bytes()
    active_before = runtime._datasets.active_info()

    assert app.update(
        "Equipment", "E-SAVE-ROLLBACK", {"status": "Inactive"}, "save-rollback-update"
    )["status"] == "COMMITTED"
    state_before = app.store.snapshot()
    original_save = app.store.save

    def fail_save():
        raise OSError("simulated dataset audit persistence failure")

    app.store.save = fail_save
    try:
        with pytest.raises(OSError, match="simulated dataset audit persistence failure"):
            runtime.save_dataset()
    finally:
        app.store.save = original_save

    assert target.read_bytes() == prior_file
    assert app.store.snapshot() == state_before
    assert runtime._datasets.active_info() == active_before
    assert not list(tmp_path.glob("*.restore-*"))


def test_runtime_dataset_save_as_removes_new_file_when_audit_persistence_fails(tmp_path):
    app = ApplicationService(PersistenceStore())
    runtime = RuntimeInterface(application=app)
    runtime._datasets.directory = tmp_path

    assert app.create("Equipment", valid_equipment("E-SAVE-AS-ROLLBACK"), "save-as-create")["status"] == "COMMITTED"
    state_before = app.store.snapshot()
    active_before = runtime._datasets.active_info()
    target = tmp_path / "Uncommitted Copy.json"
    temp = target.with_suffix(".json.tmp")
    original_save = app.store.save

    def fail_save():
        raise OSError("simulated dataset audit persistence failure")

    app.store.save = fail_save
    try:
        with pytest.raises(OSError, match="simulated dataset audit persistence failure"):
            runtime.save_dataset_as("Uncommitted Copy")
    finally:
        app.store.save = original_save

    assert not target.exists()
    assert not temp.exists()
    assert app.store.snapshot() == state_before
    assert runtime._datasets.active_info() == active_before
    assert not list(tmp_path.glob("*.restore-*"))


def test_runtime_dataset_load_replaces_runtime_data(tmp_path):
    app = ApplicationService(PersistenceStore())
    runtime = RuntimeInterface(application=app)
    runtime._datasets.directory = tmp_path

    assert app.create("Equipment", valid_equipment("E-ORIGINAL"), "dataset-original")["status"] == "COMMITTED"
    saved = runtime.save_dataset("Original State")

    assert app.create("Equipment", valid_equipment("E-SECOND"), "dataset-second")["status"] == "COMMITTED"
    assert app.store.exists("Equipment", "E-SECOND")

    loaded = runtime.load_dataset(saved["dataset_name"])
    assert loaded["status"] == "LOADED"
    assert app.store.exists("Equipment", "E-ORIGINAL")
    assert not app.store.exists("Equipment", "E-SECOND")
    assert runtime.list_datasets()["active"]["dataset_name"] == "Original State"


def test_runtime_dataset_load_failure_restores_data_lists_audit_and_active_state(tmp_path):
    app = ApplicationService(PersistenceStore())
    runtime = RuntimeInterface(application=app)
    runtime._datasets.directory = tmp_path

    assert app.create("Equipment", valid_equipment("E-ORIGINAL"), "dataset-load-original")["status"] == "COMMITTED"
    app.store.controlled_lists["equipment_category"] = {"Heavy Equipment", "Stored Category"}
    stored = runtime.save_dataset("Stored State")

    assert app.create("Equipment", valid_equipment("E-SECOND"), "dataset-load-second")["status"] == "COMMITTED"
    app.store.controlled_lists["equipment_category"] = {"Heavy Equipment", "Current Category"}
    runtime.save_dataset_as("Current State")

    before_data_and_audit = app.store.snapshot()
    before_lists = {
        key: set(values) for key, values in app.store.controlled_lists.items()
    }
    before_active = runtime._datasets.active_info()
    original_save = app.store.save

    def fail_save():
        raise OSError("simulated audit persistence failure")

    app.store.save = fail_save
    try:
        with pytest.raises(OSError, match="simulated audit persistence failure"):
            runtime.load_dataset(stored["dataset_name"])
    finally:
        app.store.save = original_save

    assert app.store.snapshot() == before_data_and_audit
    assert app.store.controlled_lists == before_lists
    assert runtime._datasets.active_info() == before_active
    assert app.store.exists("Equipment", "E-ORIGINAL")
    assert app.store.exists("Equipment", "E-SECOND")


def valid_kpi_snapshot():
    return {
        "snapshot_version": "KPI-SNAPSHOT-1",
        "scope_type": "FLEET",
        "scope_id": "FLEET:ALL",
        "period_id": "2026-10-04",
        "status": "FINAL_CANDIDATE",
        "calculation_version": "KPI-FOUNDATION-E2E-0.1",
        "policy": {"id": "PA-UA-BASELINE", "version": "1.0"},
        "time_baseline": {"id": "TB-PROJECT-DAY-0600-1800", "version": "1.0"},
        "result": {
            "PA": {"kpi": "PA", "status": "READY", "value": 81.8},
            "UA": {"kpi": "UA", "status": "READY", "value": 77.7},
            "EU": {"kpi": "EU", "status": "PENDING_DEFINITION", "value": None},
        },
        "event_versions": [],
        "population": {"total": 1, "eligible": 1, "excluded": 0},
        "contributing_hours": {"scheduledHours": 11, "availableHours": 9, "usedHours": 7},
        "exclusions": [],
    }


def test_kpi_snapshot_retry_reconciles_audit_after_persistence_failure():
    app = ApplicationService(PersistenceStore())
    runtime = RuntimeInterface(application=app)
    real_save = app.store.save
    save_attempts = {"count": 0}

    def fail_first_save():
        save_attempts["count"] += 1
        if save_attempts["count"] == 1:
            raise OSError("simulated audit persistence failure")
        return real_save()

    app.store.save = fail_first_save
    candidate = valid_kpi_snapshot()

    first = runtime.finalize_kpi(candidate, source="test")
    assert first["status"] == "COMMITTED"
    assert first["audit_status"] == "NOT_APPENDED"
    assert runtime.read("AuditLog") == []

    retry = runtime.finalize_kpi(candidate, source="test")
    assert retry["status"] == "EXISTING"
    assert retry["audit_status"] == "APPENDED"

    events = runtime.read("AuditLog")
    assert len(events) == 1
    assert events[0]["entity"] == "KPI_Snapshot"
    assert events[0]["entity_id"] == first["snapshot_id"]
    assert events[0]["action"] == "FINALIZE"
    finalized_payload = json.loads(events[0]["new_value"])
    assert finalized_payload["snapshot_id"] == first["snapshot_id"]
    assert finalized_payload["revision"] == 1
    assert finalized_payload["status"] == "FINAL"
    assert finalized_payload["content_fingerprint"]

    repeated = runtime.finalize_kpi(candidate, source="test")
    assert repeated["status"] == "EXISTING"
    assert repeated["audit_status"] == "ALREADY_APPENDED"
    assert len(runtime.read("AuditLog")) == 1
