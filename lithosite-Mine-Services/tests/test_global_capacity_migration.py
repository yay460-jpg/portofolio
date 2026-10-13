from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.mine_services import schema
from src.mine_services import migrate_global_capacity as migration_module
from src.mine_services.migrate_global_capacity import migrate


def make_partial_a3_workbook(path, version="A.3", bad_workfront_header=False, unmapped_workfront_data=False):
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet in schema.SHEETS:
        if sheet in {"GlobalCapacity", "Checker"}:
            continue
        worksheet = workbook.create_sheet(sheet)
        if sheet == "_Baseline":
            worksheet.append(["baseline_status", "LOCKED"])
        elif sheet == "_System":
            worksheet.append(["schema_version", version])
        elif sheet == "_Lists":
            worksheet.append(["measurement"])
            worksheet.append(["ton"])
        elif sheet in schema.HEADERS:
            headers = list(schema.HEADERS[sheet])
            if sheet == "Equipment":
                headers = headers[:-1]  # This is handled by the separate Equipment migration.
            elif sheet == "WorkFront":
                headers = headers[:-1]
                if bad_workfront_header:
                    headers[1] = "unexpected_field"
                worksheet.append(headers)
                worksheet.append([
                    "WF-001", "Road & Hauling", "Pit A", "Operations", "Active", None, None
                ])
                if unmapped_workfront_data:
                    worksheet.cell(row=2, column=len(headers) + 1, value="preserve-me")
                continue
            elif sheet == "Operations":
                headers = headers[:15]  # Known A.3 prefix before hauling snapshot fields.
            worksheet.append(headers)
            if sheet == "Operations":
                worksheet.append([
                    "OPS-001", "2026-10-12", "06:00", "Road & Hauling", "WF-001",
                    "EQ-001", "Hauling", 110, "ton", 4, 4, "DRAFT", "test", None, None
                ])
    workbook.save(path)


def test_global_capacity_migration_repairs_known_a3_prefixes_and_preserves_rows(tmp_path):
    path = tmp_path / "partial-a3.xlsx"
    make_partial_a3_workbook(path)

    assert migrate(path) == path

    workbook = load_workbook(path, data_only=True)
    try:
        assert list(workbook.sheetnames) == schema.SHEETS
        for sheet in ("GlobalCapacity", "Checker", "WorkFront", "Operations"):
            assert [cell.value for cell in workbook[sheet][1]] == schema.HEADERS[sheet]
        assert workbook["_System"]["B1"].value == "A.3"
        workfront = list(workbook["WorkFront"].iter_rows(min_row=2, values_only=True))[0]
        operation = list(workbook["Operations"].iter_rows(min_row=2, values_only=True))[0]
        assert workfront[:7] == ("WF-001", "Road & Hauling", "Pit A", "Operations", "Active", None, None)
        assert workfront[7] is None
        assert operation[:15] == (
            "OPS-001", "2026-10-12", "06:00", "Road & Hauling", "WF-001",
            "EQ-001", "Hauling", 110, "ton", 4, 4, "DRAFT", "test", None, None
        )
        assert operation[15:] == (None,) * (len(schema.HEADERS["Operations"]) - 15)
    finally:
        workbook.close()


def test_global_capacity_migration_rejects_same_length_wrong_header_without_mutation(tmp_path):
    path = tmp_path / "bad-header-a3.xlsx"
    make_partial_a3_workbook(path, bad_workfront_header=True)
    original = path.read_bytes()

    with pytest.raises(ValueError, match="HEADER_MISMATCH:WorkFront"):
        migrate(path)

    assert path.read_bytes() == original
    assert list(tmp_path.glob("bad-header-a3.*.xlsx")) == []


def test_global_capacity_migration_rejects_wrong_schema_version_without_mutation(tmp_path):
    path = tmp_path / "wrong-version.xlsx"
    make_partial_a3_workbook(path, version="A.2")
    original = path.read_bytes()

    with pytest.raises(ValueError, match="SCHEMA_VERSION"):
        migrate(path)

    assert path.read_bytes() == original
    assert list(tmp_path.glob("wrong-version.*.xlsx")) == []


def test_global_capacity_migration_rejects_unmapped_row_data_without_mutation(tmp_path):
    path = tmp_path / "unmapped-a3.xlsx"
    make_partial_a3_workbook(path, unmapped_workfront_data=True)
    original = path.read_bytes()

    with pytest.raises(ValueError, match="UNMAPPED_COLUMN_DATA:WorkFront"):
        migrate(path)

    assert path.read_bytes() == original
    assert list(tmp_path.glob("unmapped-a3.*.xlsx")) == []


def test_global_capacity_migration_preserves_source_if_atomic_replace_fails(tmp_path, monkeypatch):
    path = tmp_path / "replace-failure-a3.xlsx"
    make_partial_a3_workbook(path)
    original = path.read_bytes()
    real_replace = Path.replace

    def fail_target_replace(source, target):
        if Path(target) == path and Path(source) != path:
            raise OSError("simulated GlobalCapacity replacement failure")
        return real_replace(source, target)

    monkeypatch.setattr(Path, "replace", fail_target_replace)
    with pytest.raises(OSError, match="simulated GlobalCapacity replacement failure"):
        migrate(path)

    assert path.read_bytes() == original
    assert list(tmp_path.glob("replace-failure-a3.*.xlsx")) == []
