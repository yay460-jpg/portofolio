from openpyxl import Workbook, load_workbook

from src.mine_services.application import ApplicationService
from src.mine_services.persistence import PersistenceStore
from src.mine_services import schema


def service():
    return ApplicationService(store=PersistenceStore(schema_module=schema))


def create(svc, entity, row, request_id):
    result = svc.create(entity, row, request_id)
    assert result["status"] == "COMMITTED", result
    return result


def test_equipment_capacity_profile_is_canonical_fk():
    assert "capacity_profile_id" in schema.HEADERS["Equipment"]
    assert schema.FK["Equipment.capacity_profile_id"] == ("GlobalCapacity", "capacity_profile_id", False)


def test_operation_uses_equipment_capacity_not_workfront_capacity():
    svc = service()
    create(svc, "GlobalCapacity", {
        "capacity_profile_id": "GC-EQ",
        "capacity_name": "Equipment Capacity",
        "capacity_value": 27.5,
        "measurement": "ton",
        "status": "Active",
    }, "gc-eq")
    create(svc, "GlobalCapacity", {
        "capacity_profile_id": "GC-WF",
        "capacity_name": "Legacy Work Front Capacity",
        "capacity_value": 35.5,
        "measurement": "ton",
        "status": "Active",
    }, "gc-wf")
    create(svc, "WorkFront", {
        "work_front_id": "WF-A",
        "domain": "Road & Hauling",
        "location": "Pit A",
        "responsible": "Ops",
        "status": "Active",
        "capacity_profile_id": "GC-WF",
    }, "wf")
    create(svc, "Equipment", {
        "equipment_id": "DT-001",
        "unit_no": "DT-001",
        "category": "Heavy Equipment",
        "type": "Dump Truck",
        "owner_type": "Owner",
        "owner_name": "",
        "status": "Active",
        "effective_from": "2026-10-09",
        "effective_to": None,
        "capacity_profile_id": "GC-EQ",
    }, "eq")
    create(svc, "Operations", {
        "transaction_id": "OPS-EQ-CAP",
        "transaction_date": "2026-10-09",
        "transaction_time": "06:00",
        "domain": "Road & Hauling",
        "work_front_id": "WF-A",
        "equipment_id": "DT-001",
        "activity": "Hauling",
        "retase": 4,
        "measurement": "",
        "actual_hours": 4,
        "target_hours": 4,
        "status": "DRAFT",
        "source": "Checker",
    }, "op")
    row = svc.read("Operations", "OPS-EQ-CAP")
    assert row["capacity_profile_id"] == "GC-EQ"
    assert row["applied_capacity"] == 27.5
    assert row["quantity"] == 110.0


def test_dump_truck_hauling_rejected_without_equipment_capacity_profile():
    svc = service()
    create(svc, "WorkFront", {
        "work_front_id": "WF-A",
        "domain": "Road & Hauling",
        "location": "Pit A",
        "responsible": "Ops",
        "status": "Active",
        "capacity_profile_id": None,
    }, "wf")
    create(svc, "Equipment", {
        "equipment_id": "DT-001",
        "unit_no": "DT-001",
        "category": "Heavy Equipment",
        "type": "Dump Truck",
        "owner_type": "Owner",
        "owner_name": "",
        "status": "Active",
        "effective_from": "2026-10-09",
        "effective_to": None,
        "capacity_profile_id": None,
    }, "eq")
    result = svc.create("Operations", {
        "transaction_id": "OPS-NO-CAP",
        "transaction_date": "2026-10-09",
        "transaction_time": "06:00",
        "domain": "Road & Hauling",
        "work_front_id": "WF-A",
        "equipment_id": "DT-001",
        "activity": "Hauling",
        "retase": 4,
        "measurement": "",
        "actual_hours": 4,
        "target_hours": 4,
        "status": "DRAFT",
        "source": "Checker",
    }, "op")
    assert result["status"] == "REJECTED"
    assert any(error.field == "capacity_profile_id" for error in result["errors"])


def test_global_capacity_delete_is_blocked_by_equipment_reference():
    svc = service()
    create(svc, "GlobalCapacity", {
        "capacity_profile_id": "GC-REF",
        "capacity_name": "Referenced",
        "capacity_value": 27.5,
        "measurement": "ton",
        "status": "Active",
    }, "gc")
    create(svc, "Equipment", {
        "equipment_id": "DT-REF",
        "unit_no": "DT-REF",
        "category": "Heavy Equipment",
        "type": "Dump Truck",
        "owner_type": "Owner",
        "owner_name": "",
        "status": "Active",
        "effective_from": "2026-10-09",
        "effective_to": None,
        "capacity_profile_id": "GC-REF",
    }, "eq")
    result = svc.delete("GlobalCapacity", "GC-REF", "gc-delete")
    assert result["status"] == "REJECTED"
    assert any(error["code"] == "APP-004" for error in result["errors"])


def test_a3_equipment_header_migration_appends_capacity_profile(tmp_path):
    path = tmp_path / "equipment-capacity-migration.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    baseline = wb.create_sheet("_Baseline")
    baseline.append(["baseline_status", "LOCKED"])
    system = wb.create_sheet("_System")
    system.append(["schema_version", schema.SCHEMA_VERSION])
    lists = wb.create_sheet("_Lists")
    defaults = PersistenceStore().controlled_lists
    headers = list(defaults)
    lists.append(headers)
    for i in range(max(len(values) for values in defaults.values())):
        lists.append([sorted(defaults[h])[i] if i < len(defaults[h]) else None for h in headers])

    for entity in schema.DOMAIN_ENTITIES:
        ws = wb.create_sheet(entity)
        headers_for_entity = list(schema.HEADERS[entity])
        if entity == "Equipment":
            headers_for_entity = headers_for_entity[:-1]
        ws.append(headers_for_entity)

    audit = wb.create_sheet("AuditLog")
    audit.append(schema.HEADERS["AuditLog"])
    wb.save(path)

    store = PersistenceStore(path=path, schema_module=schema)
    assert store.schema.HEADERS["Equipment"][-1] == "capacity_profile_id"

    check = load_workbook(path, data_only=True)
    assert check["Equipment"][1].value == "equipment_id"
    assert check["Equipment"][1].value is not None
    assert check["Equipment"].max_column == len(schema.HEADERS["Equipment"])


def test_v39_conservative_migration_resolves_unique_operation_profile(tmp_path):
    from src.mine_services.migrate_v39_equipment_capacity import migrate

    path = tmp_path / "equipment-capacity-v39.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    baseline = wb.create_sheet("_Baseline")
    baseline.append(["baseline_status", "LOCKED"])
    system = wb.create_sheet("_System")
    system.append(["schema_version", schema.SCHEMA_VERSION])
    lists = wb.create_sheet("_Lists")
    defaults = PersistenceStore().controlled_lists
    headers = list(defaults)
    lists.append(headers)
    for i in range(max(len(values) for values in defaults.values())):
        lists.append([sorted(defaults[h])[i] if i < len(defaults[h]) else None for h in headers])

    for entity in schema.DOMAIN_ENTITIES:
        wb.create_sheet(entity).append(schema.HEADERS[entity])

    wb["GlobalCapacity"].append(["GC-275", "10 Wheel Heavy", "Volvo", 27.5, "ton", "Active", None, None])
    wb["WorkFront"].append(["WF-A", "Road & Hauling", "Pit A", "Ops", "Active", None, None, "GC-275"])
    wb["Equipment"].append(["DT-001", "DT-001", "Heavy Equipment", "Dump Truck", "Owner", "", "Active", "2026-10-09", None, None])
    wb["Operations"].append([
        "OPS-001", "2026-10-09", "06:00", "Road & Hauling", "WF-A", "DT-001",
        "Hauling", 110, "ton", 4, 4, "DRAFT", "Checker", None, None, 4,
        None, 27.5, "ton", "10:00", "Day", "Ore", "Field Checker"
    ])
    wb["AuditLog"].append(schema.HEADERS["AuditLog"])
    wb.save(path)

    resolved, ambiguous, untouched = migrate(path)
    assert resolved == [("DT-001", "GC-275")]
    assert ambiguous == []
    assert untouched == []

    migrated = load_workbook(path, data_only=True)
    equipment_row = list(migrated["Equipment"].iter_rows(min_row=2, values_only=True))[0]
    assert equipment_row[-1] == "GC-275"
