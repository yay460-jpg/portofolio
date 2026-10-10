"""Build an isolated, deterministic A3 workbook for CI-only database regression tests.

The checked-in operational workbook is not used or replaced by this fixture. Local
runs keep using the developer's configured/working database unless
MINE_SERVICES_TEST_DB is explicitly supplied.
"""
import os
from pathlib import Path

from openpyxl import Workbook


def main() -> None:
    destination = Path(os.environ["MINE_SERVICES_TEST_DB"])
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    system = workbook.active
    system.title = "_System"
    system.append(["key", "value"])
    system.append(["schema_version", "A.3"])

    lists = workbook.create_sheet("_Lists")
    lists.append(["service_domain", "measurement", "equipment_category", "equipment_type"])
    lists.append(["Mining", "cycle", "Heavy Equipment", "Dump Truck"])

    equipment = workbook.create_sheet("Equipment")
    equipment.append([
        "equipment_id", "unit_no", "category", "type", "owner_type", "owner_name",
        "status", "effective_from", "effective_to", "capacity_profile_id",
    ])
    equipment.append([
        "DT-CI-001", "DT-CI-001", "Heavy Equipment", "Dump Truck", "Contractor",
        "CI Fixture", "Active", "2026-10-01", None, None,
    ])

    operations = workbook.create_sheet("Operations")
    operations.append([
        "transaction_id", "transaction_date", "transaction_time", "domain",
        "work_front_id", "equipment_id", "activity", "quantity", "measurement",
        "actual_hours", "target_hours", "status", "source", "created_at",
        "updated_at", "retase", "capacity_profile_id", "applied_capacity",
        "capacity_measurement", "end_time", "shift", "material", "checker_name",
    ])

    maintenance = workbook.create_sheet("Maintenance")
    maintenance.append([
        "maintenance_id", "equipment_id", "event_date", "event_type", "failure_code",
        "start_time", "end_time", "downtime_hours", "action", "status", "source",
    ])

    workbook.save(destination)
    workbook.close()
    print(f"Created isolated CI workbook fixture: {destination}")


if __name__ == "__main__":
    main()
