from pathlib import Path
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))

from mine_services import PersistenceStore
from mine_services import schema as schema_a2
from mine_services import schema_a3
from mine_services.schema_migration import migrate_a2_to_a3


def _make_a2_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "_Baseline"
    ws.append(["baseline_id"])
    ws.append(["BASE-01"])

    system = wb.create_sheet("_System")
    system.append(["key", "value"])
    system.append(["schema_version", "A.2"])

    lists = wb.create_sheet("_Lists")
    lists.append(["service_domain"])
    lists.append(["Road & Hauling"])

    for entity in schema_a2.DOMAIN_ENTITIES:
        ws = wb.create_sheet(entity)
        ws.append(schema_a2.HEADERS[entity])

    audit = wb.create_sheet("AuditLog")
    audit.append(schema_a2.HEADERS["AuditLog"])
    wb.save(path)


def test_kf_a2_remains_default_schema():
    assert schema_a2.SCHEMA_VERSION == "A.2"
    assert "MapMarker" not in schema_a2.DOMAIN_ENTITIES
    assert "MapMarker" not in schema_a2.HEADERS


def test_kf_a3_is_explicit_opt_in():
    assert schema_a3.SCHEMA_VERSION == "A.3"
    assert "MapMarker" in schema_a3.DOMAIN_ENTITIES
    assert schema_a3.PKS["MapMarker"] == "marker_id"


def test_kf_migration_does_not_modify_source(tmp_path):
    source = tmp_path / "source-a2.xlsx"
    target = tmp_path / "target-a3.xlsx"
    _make_a2_workbook(source)
    source_before = source.read_bytes()

    migrate_a2_to_a3(source, target)

    assert source.read_bytes() == source_before
    assert target.exists()


def test_kf_migrated_a3_loads_with_a3_persistence(tmp_path):
    source = tmp_path / "source-a2.xlsx"
    target = tmp_path / "target-a3.xlsx"
    _make_a2_workbook(source)
    migrate_a2_to_a3(source, target)

    store = PersistenceStore(target, schema_module=schema_a3)
    assert store.schema.SCHEMA_VERSION == "A.3"
    assert store.all("MapMarker") == []

    marker = {
        "marker_id": "MK-KF-001",
        "marker_type": "HSE",
        "label": "HSE-KF",
        "easting": 450000.0,
        "northing": 9700000.0,
        "elevation": 25.0,
        "source_entity": "HSE",
        "source_id": "HSE-KF-001",
        "status": "ACTIVE",
    }
    store.insert("MapMarker", marker["marker_id"], marker)
    store.save()

    reloaded = PersistenceStore(target, schema_module=schema_a3)
    assert reloaded.get("MapMarker", "MK-KF-001") == marker


def test_kf_a2_store_rejects_a3_workbook(tmp_path):
    source = tmp_path / "source-a2.xlsx"
    target = tmp_path / "target-a3.xlsx"
    _make_a2_workbook(source)
    migrate_a2_to_a3(source, target)

    try:
        PersistenceStore(target)
    except ValueError as exc:
        assert "SCHEMA_VERSION" in str(exc)
    else:
        raise AssertionError("A.2 default store must not silently open A.3 workbook")


def test_kf_a3_marker_sheet_is_not_added_to_hse_headers():
    assert schema_a3.HEADERS["HSE"] == schema_a2.HEADERS["HSE"]
    assert "easting" not in schema_a3.HEADERS["HSE"]
    assert "northing" not in schema_a3.HEADERS["HSE"]
    assert "elevation" not in schema_a3.HEADERS["HSE"]
