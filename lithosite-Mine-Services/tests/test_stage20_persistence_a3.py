import sys
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))

from mine_services.persistence import PersistenceStore
from mine_services.schema import SCHEMA_VERSION, SHEETS, HEADERS, CONTROLLED, NIK_HEADERS
from mine_services.application import ApplicationService
from mine_services.runtime import RuntimeInterface
from mine_services.adapter import RuntimeAdapter


def make_a3_workbook(path: Path):
    wb = Workbook()
    wb.remove(wb.active)

    for sheet in SHEETS:
        ws = wb.create_sheet(sheet)
        if sheet in HEADERS:
            ws.append(HEADERS[sheet])
        elif sheet == "NIK":
            from mine_services.schema import NIK_HEADERS
            ws.append(NIK_HEADERS)
        elif sheet == "_System":
            ws.append(["schema_version", "A.3"])
        elif sheet == "_Lists":
            list_names = list(dict.fromkeys(CONTROLLED.values()))
            ws.append(list_names)
            for row_index in range(max(len(PersistenceStore().controlled_lists[name]) for name in list_names)):
                ws.append([
                    sorted(PersistenceStore().controlled_lists[name])[row_index]
                    if row_index < len(PersistenceStore().controlled_lists[name])
                    else None
                    for name in list_names
                ])
        elif sheet == "_Baseline":
            ws.append(["baseline"])
            ws.append(["A.3 baseline"])

    wb.save(path)


def test_stage20_a3_persistence_store_is_default():
    store = PersistenceStore()
    assert store.schema.SCHEMA_VERSION == "A.3"
    assert "MapMarker" in store.schema.DOMAIN_ENTITIES
    assert SCHEMA_VERSION == "A.3"


def test_stage20_a3_persistence_store_holds_mapmarker_in_memory():
    store = PersistenceStore()
    row = {
        "marker_id": "MK-A3-001",
        "marker_type": "HSE",
        "label": "HSE-001",
        "easting": 450000.0,
        "northing": 9700000.0,
        "elevation": 25.5,
        "source_entity": "HSE",
        "source_id": "HSE-001",
        "status": "ACTIVE",
    }
    store.insert("MapMarker", row["marker_id"], row)

    assert store.get("MapMarker", "MK-A3-001") == row
    assert store.all("MapMarker") == [row]


def test_stage20_a3_persistence_store_roundtrips_canonical_workbook(tmp_path):
    target = tmp_path / "a3.xlsx"
    make_a3_workbook(target)

    store = PersistenceStore(target)
    assert store.schema.SCHEMA_VERSION == "A.3"
    assert store.get("MapMarker", "missing") is None

    row = {
        "marker_id": "MK-A3-002",
        "marker_type": "WORKFRONT",
        "label": "WF-002",
        "easting": 451000.0,
        "northing": 9701000.0,
        "elevation": 30.0,
        "source_entity": "WorkFront",
        "source_id": "WF-002",
        "status": "ACTIVE",
    }
    store.insert("MapMarker", row["marker_id"], row)
    store.save()

    saved = load_workbook(target, data_only=True)
    marker_rows = list(saved["MapMarker"].values)
    assert marker_rows[1] == tuple(
        row[field] for field in store.schema.HEADERS["MapMarker"]
    )


def test_stage20_a3_persistence_store_rejects_non_a3_workbook(tmp_path):
    target = tmp_path / "legacy.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "_System"
    ws.append(["schema_version", "A.2"])
    wb.save(target)

    try:
        PersistenceStore(target)
    except ValueError as exc:
        assert "SCHEMA_VERSION" in str(exc)
    else:
        raise AssertionError("A.3 default store must reject a non-A3 workbook")


def test_stage20_a3_persistence_migrates_top_soil_into_checker_material_list(tmp_path):
    target = tmp_path / "a3-top-soil.xlsx"
    make_a3_workbook(target)

    workbook = load_workbook(target)
    lists = workbook["_Lists"]
    headers = [cell.value for cell in lists[1]]
    material_column = headers.index("checker_material") + 1
    for row_index in range(2, lists.max_row + 1):
        if lists.cell(row=row_index, column=material_column).value == "Top Soil":
            lists.cell(row=row_index, column=material_column).value = None
    workbook.save(target)

    first_store = PersistenceStore(target)
    assert "Top Soil" in first_store.controlled_lists["checker_material"]

    # Reopening should be idempotent: do not add duplicate vocabulary entries.
    PersistenceStore(target)
    saved = load_workbook(target, data_only=True)
    stored_materials = [
        saved["_Lists"].cell(row=row_index, column=material_column).value
        for row_index in range(2, saved["_Lists"].max_row + 1)
    ]
    assert stored_materials.count("Top Soil") == 1


def test_stage20_a3_persistence_removes_legacy_plans_target_hours_without_shifting_values(tmp_path):
    target = tmp_path / "a3-plan-hours.xlsx"
    make_a3_workbook(target)

    workbook = load_workbook(target)
    plans = workbook["Plans"]
    # Simulate a previously valid A.3 workbook which still has the retired column.
    plans.insert_cols(10, 1)
    plans.cell(row=1, column=10, value="target_hours")
    plans.auto_filter.ref = "A1:K2"
    old_row = {
        "plan_id": "PLN-MIG-001",
        "period": "2026-10",
        "start_date": "2026-10-01",
        "end_date": "2026-10-31",
        "domain": "Road & Hauling",
        "work_front_id": "WF-MIG-001",
        "activity": "Hauling",
        "target_quantity": 500,
        "measurement": "ton",
        "target_hours": 8,
        "status": "Draft",
    }
    headers = [cell.value for cell in plans[1]]
    for column, header in enumerate(headers, start=1):
        plans.cell(row=2, column=column, value=old_row.get(header))
    workbook.save(target)

    store = PersistenceStore(target)
    migrated = load_workbook(target, data_only=True)
    migrated_headers = [cell.value for cell in migrated["Plans"][1]]

    assert migrated_headers == HEADERS["Plans"]
    assert "target_hours" not in migrated_headers
    assert migrated["Plans"].auto_filter.ref == "A1:J2"
    row = store.get("Plans", "PLN-MIG-001")
    assert row["target_quantity"] == 500
    assert row["measurement"] == "ton"
    assert row["status"] == "Draft"
    assert row["start_date"] == "2026-10-01"
    assert row["end_date"] == "2026-10-31"
    assert "target_hours" not in row



def test_stage20_a3_persistence_updates_plans_excel_table_after_removing_target_hours(tmp_path):
    target = tmp_path / "a3-plan-table.xlsx"
    make_a3_workbook(target)

    workbook = load_workbook(target)
    plans = workbook["Plans"]
    plans.insert_cols(10, 1)
    plans.cell(row=1, column=10, value="target_hours")
    old_row = {
        "plan_id": "PLN-TABLE-001",
        "period": "2026-10",
        "start_date": "2026-10-01",
        "end_date": "2026-10-31",
        "domain": "Road & Hauling",
        "work_front_id": "WF-MIG-001",
        "activity": "Hauling",
        "target_quantity": 500,
        "measurement": "ton",
        "target_hours": 8,
        "status": "Draft",
    }
    headers = [cell.value for cell in plans[1]]
    for column, header in enumerate(headers, start=1):
        plans.cell(row=2, column=column, value=old_row.get(header))
    plans.add_table(Table(displayName="PlansTable", ref="A1:K2"))
    workbook.save(target)

    PersistenceStore(target)
    migrated = load_workbook(target, data_only=True)
    table = migrated["Plans"].tables["PlansTable"]

    assert table.ref == "A1:J2"
    assert len(table.tableColumns) == len(HEADERS["Plans"])
    assert all(column.name != "target_hours" for column in table.tableColumns)




def test_stage20_a3_persistence_adds_nik_headers_idempotently_and_preserves_records(tmp_path):
    target = tmp_path / "a3-without-nik.xlsx"
    make_a3_workbook(target)

    workbook = load_workbook(target)
    workbook.remove(workbook["NIK"])
    workbook["Equipment"].append([
        "EQ-NIK-PRESERVE", "DT-NIK-PRESERVE", "Heavy Equipment", "Dump Truck",
        "Owner", "Preserve Existing", "Active", "2026-10-15", None, None,
    ])
    workbook.save(target)

    first_store = PersistenceStore(target)
    assert first_store.get("Equipment", "EQ-NIK-PRESERVE")["unit_no"] == "DT-NIK-PRESERVE"

    migrated = load_workbook(target, data_only=True)
    from mine_services.schema import NIK_HEADERS
    assert "NIK" in migrated.sheetnames
    assert migrated.sheetnames.index("NIK") == migrated.sheetnames.index("Equipment") + 1
    assert list(migrated["NIK"].values) == [tuple(NIK_HEADERS)]
    assert migrated["Equipment"]["A2"].value == "EQ-NIK-PRESERVE"

    # Reopening must not add another sheet or duplicate the headers.
    PersistenceStore(target)
    reopened = load_workbook(target, data_only=True)
    assert reopened.sheetnames.count("NIK") == 1
    assert list(reopened["NIK"].values) == [tuple(NIK_HEADERS)]


def test_stage20_a3_persistence_upgrades_empty_nik_placeholder_in_place(tmp_path):
    target = tmp_path / "a3-empty-nik-placeholder.xlsx"
    make_a3_workbook(target)

    workbook = load_workbook(target)
    nik = workbook["NIK"]
    for column in range(1, 6):
        nik.cell(row=1, column=column, value=None)
    workbook.save(target)

    PersistenceStore(target)
    migrated = load_workbook(target, data_only=True)
    from mine_services.schema import NIK_HEADERS
    assert list(migrated["NIK"].values) == [tuple(NIK_HEADERS)]
    assert migrated.sheetnames.count("NIK") == 1

def test_stage20_nik_create_appends_to_canonical_sheet_and_is_readable(tmp_path):
    target = tmp_path / "a3-nik-create.xlsx"
    make_a3_workbook(target)
    store = PersistenceStore(target)
    runtime = RuntimeInterface(application=ApplicationService(store=store))
    adapter = RuntimeAdapter(runtime)
    response = adapter.handle({
        "request_id": "create-nik-001",
        "operation": "CREATE",
        "entity": "NIK",
        "row": {
            "Effective From": "2026-10-15",
            "Effective To": "",
            "Id No (NIK)": "320101001",
            "Origin": "Makassar",
            "Name Employment": "Test Employee",
            "Information": "Mine Services",
        },
    })
    assert response["status"] == "COMMITTED"
    assert response["entity_id"] == "320101001"
    records = store.all_nik_records()
    assert len(records) == 1
    assert records[0]["Id No (NIK)"] == "320101001"
    assert records[0]["Name Employment"] == "Test Employee"
    reopened = load_workbook(target, data_only=True)
    assert reopened["NIK"].max_row == 2
    assert reopened["Equipment"].max_row == 1


def test_stage20_nik_create_rejects_missing_identity_fields(tmp_path):
    target = tmp_path / "a3-nik-required.xlsx"
    make_a3_workbook(target)
    store = PersistenceStore(target)
    try:
        store.create_nik_record({"Name Employment": "Missing NIK"})
    except ValueError as error:
        assert "Id No (NIK) and Name Employment are required" in str(error)
    else:
        raise AssertionError("NIK identity fields must be required")


def test_stage20_nik_create_rejects_duplicate_nik_without_overwriting_existing_record(tmp_path):
    target = tmp_path / "a3-nik-duplicate.xlsx"
    make_a3_workbook(target)
    workbook = load_workbook(target)
    workbook["NIK"].append(["2026-01-01", None, "320101001", "Makassar", "Existing Employee", ""])
    workbook.save(target)
    store = PersistenceStore(target)
    try:
        store.create_nik_record({"Id No (NIK)": "320101001", "Name Employment": "Duplicate Employee"})
    except ValueError as error:
        assert "DUPLICATE_NIK" in str(error)
    else:
        raise AssertionError("Duplicate NIK must be rejected")
    records = store.all_nik_records()
    assert len(records) == 1
    assert records[0]["Name Employment"] == "Existing Employee"


def test_stage20_nik_list_reads_canonical_rows_through_runtime_adapter(tmp_path):
    target = tmp_path / "a3-nik-list.xlsx"
    make_a3_workbook(target)
    workbook = load_workbook(target)
    workbook["NIK"].append([
        "2026-01-12", "2026-12-31", "NIK-001", "Sulawesi",
        "Mine Services", "Permanent staff",
    ])
    workbook["NIK"].append([None, None, None, None, None, None])
    workbook.save(target)

    store = PersistenceStore(target)
    runtime = RuntimeInterface(application=ApplicationService(store=store))
    adapter = RuntimeAdapter(runtime)
    response = adapter.handle({
        "request_id": "read-nik-list-001",
        "operation": "READ",
        "entity": "NIK",
    })

    assert response["status"] == "OK"
    assert response["data"] == [{
        NIK_HEADERS[0]: "2026-01-12",
        NIK_HEADERS[1]: "2026-12-31",
        NIK_HEADERS[2]: "NIK-001",
        NIK_HEADERS[3]: "Sulawesi",
        NIK_HEADERS[4]: "Mine Services",
        NIK_HEADERS[5]: "Permanent staff",
    }]
    assert "NIK" not in store.schema.DOMAIN_ENTITIES
    assert "NIK" in RuntimeAdapter.READ_ENTITIES

def _valid_nik_record(nik_id="NIK-EVO2-01"):
    return {
        "Effective From": "2026-01-01",
        "Effective To": "2026-12-31",
        "Id No (NIK)": nik_id,
        "Origin": "Indonesia",
        "Name Employment": "EVO2 Regression Employee",
        "Information": "EVO2 test",
    }


def test_runtime_nik_crud_enforces_request_guard_and_writes_audit(tmp_path):
    target = tmp_path / "nik-runtime-contract.xlsx"
    make_a3_workbook(target)
    app = ApplicationService(PersistenceStore(target))
    runtime = RuntimeInterface(application=app)

    created = runtime.create("NIK", _valid_nik_record(), "nik-evo2-create")
    assert created == {"status": "COMMITTED", "entity_id": "NIK-EVO2-01"}
    duplicate = runtime.create("NIK", _valid_nik_record("NIK-EVO2-02"), "nik-evo2-create")
    assert duplicate == {"status": "DUPLICATE_REQUEST"}
    assert [row["Id No (NIK)"] for row in runtime.read("NIK")] == ["NIK-EVO2-01"]

    updated_record = _valid_nik_record()
    updated_record["Information"] = "updated"
    updated = runtime.update("NIK", "NIK-EVO2-01", updated_record, "nik-evo2-update")
    assert updated == {"status": "COMMITTED", "entity_id": "NIK-EVO2-01"}
    deleted = runtime.delete("NIK", "NIK-EVO2-01", "nik-evo2-delete")
    assert deleted == {"status": "COMMITTED", "entity_id": "NIK-EVO2-01"}

    audit = app.store.audit()
    assert [(event["entity"], event["action"], event["source"]) for event in audit] == [
        ("NIK", "CREATE", "nik-evo2-create"),
        ("NIK", "UPDATE", "nik-evo2-update"),
        ("NIK", "DELETE", "nik-evo2-delete"),
    ]
    saved = load_workbook(target, data_only=True)
    assert [row[2] for row in saved["NIK"].iter_rows(min_row=2, values_only=True)] == []
    assert saved["AuditLog"].max_row == 4
    saved.close()


def test_runtime_nik_audit_failure_rolls_back_workbook_and_audit(tmp_path, monkeypatch):
    target = tmp_path / "nik-runtime-rollback.xlsx"
    make_a3_workbook(target)
    app = ApplicationService(PersistenceStore(target))
    runtime = RuntimeInterface(application=app)
    original_bytes = target.read_bytes()

    def fail_audit(**kwargs):
        raise RuntimeError("injected audit failure")

    monkeypatch.setattr(app.audit_repository, "append", fail_audit)
    result = runtime.create("NIK", _valid_nik_record(), "nik-evo2-audit-fail")

    assert result["status"] == "REJECTED"
    assert "injected audit failure" in result["errors"][0]["message"]
    assert runtime.read("NIK") == []
    assert app.store.audit() == []
    assert target.read_bytes() == original_bytes



def test_runtime_nik_snapshot_read_failure_returns_rejected_without_mutation(tmp_path, monkeypatch):
    target = tmp_path / "nik-runtime-snapshot-read-failure.xlsx"
    make_a3_workbook(target)
    app = ApplicationService(PersistenceStore(target))
    runtime = RuntimeInterface(application=app)
    original_bytes = target.read_bytes()
    original_store = app.store.snapshot()
    original_read_bytes = Path.read_bytes

    def fail_target_read(path_obj):
        if path_obj == target:
            raise OSError("injected workbook snapshot read failure")
        return original_read_bytes(path_obj)

    monkeypatch.setattr(Path, "read_bytes", fail_target_read)
    result = runtime.create("NIK", _valid_nik_record(), "nik-evo2-snapshot-read-fail")

    assert result["status"] == "REJECTED"
    assert result["errors"][0]["code"] == "APP-006"
    assert "injected workbook snapshot read failure" in result["errors"][0]["message"]
    assert app.store.snapshot() == original_store
    assert runtime.read("NIK") == []
    assert app.store.audit() == []
    assert original_read_bytes(target) == original_bytes


def test_runtime_import_xlsx_unreadable_path_is_sanitized_and_non_mutating(tmp_path):
    target = tmp_path / "import-canonical.xlsx"
    make_a3_workbook(target)
    app = ApplicationService(PersistenceStore(target))
    runtime = RuntimeInterface(application=app)
    before_bytes = target.read_bytes()
    before_store = app.store.snapshot()
    missing = tmp_path / "private-path-that-must-not-leak.xlsx"

    result = runtime.import_xlsx(missing)

    assert result["status"] == "REJECTED"
    assert result["lifecycle"] == "REJECTED"
    assert result["rows_committed"] == 0
    assert result["errors"][0]["code"] == "E001"
    assert str(missing) not in result["errors"][0]["message"]
    assert "could not be opened" in result["errors"][0]["message"].lower()
    assert app.store.snapshot() == before_store
    assert app.store.audit() == []
    assert target.read_bytes() == before_bytes
