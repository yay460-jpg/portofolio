from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.mine_services import migrate_measurement as measurement_migration
from src.mine_services.migrate_measurement import migrate


def make_workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "_Lists"
    ws.append(["unit"])
    ws.append(["ton"])
    for sheet, headers in {
        "Operations": ["transaction_id", "quantity", "unit", "capacity_unit"],
        "Plans": ["plan_id", "target_quantity", "unit"],
        "GlobalCapacity": ["capacity_profile_id", "capacity_value", "unit"],
    }.items():
        ws = wb.create_sheet(sheet)
        ws.append(headers)
        if sheet == "Operations":
            ws.append(["OPS-1", 100, "ton", "ton"])
        elif sheet == "Plans":
            ws.append(["PLN-1", 100, "ton"])
        else:
            ws.append(["GC-1", 25, "ton"])
    wb.save(path)


def test_measurement_migration_renames_only_measurement_headers(tmp_path):
    path = tmp_path / "Mine-Services-Database-A3.xlsx"
    make_workbook(path)

    migrate(path)

    wb = load_workbook(path, data_only=False)
    assert [c.value for c in wb["Operations"][1]] == [
        "transaction_id", "quantity", "measurement", "capacity_measurement"
    ]
    assert [c.value for c in wb["Plans"][1]] == [
        "plan_id", "target_quantity", "measurement"
    ]
    assert [c.value for c in wb["GlobalCapacity"][1]] == [
        "capacity_profile_id", "capacity_value", "measurement"
    ]
    assert wb["_Lists"]["A1"].value == "measurement"
    assert wb["_Lists"]["A2"].value == "ton"


def test_measurement_migration_is_idempotent(tmp_path):
    path = tmp_path / "Mine-Services-Database-A3.xlsx"
    make_workbook(path)

    migrate(path)
    migrate(path)

    wb = load_workbook(path, data_only=False)
    assert wb["Operations"]["C1"].value == "measurement"
    assert wb["Operations"]["D1"].value == "capacity_measurement"


def test_measurement_migration_preserves_source_when_temp_verification_fails(tmp_path, monkeypatch):
    path = tmp_path / "Mine-Services-Database-A3.xlsx"
    make_workbook(path)
    original = path.read_bytes()
    real_verify = measurement_migration._verify_measurement_headers

    def fail_temp_verification(candidate):
        if Path(candidate) != path:
            raise RuntimeError("simulated temporary workbook verification failure")
        return real_verify(candidate)

    monkeypatch.setattr(measurement_migration, "_verify_measurement_headers", fail_temp_verification)
    with pytest.raises(RuntimeError, match="simulated temporary workbook verification failure"):
        migrate(path)

    assert path.read_bytes() == original
    assert list(tmp_path.glob("Mine-Services-Database-A3.*.xlsx")) == []


def test_measurement_migration_preserves_source_when_atomic_replace_fails(tmp_path, monkeypatch):
    path = tmp_path / "Mine-Services-Database-A3.xlsx"
    make_workbook(path)
    original = path.read_bytes()
    real_replace = Path.replace

    def fail_target_replace(source, target):
        if Path(target) == path and Path(source) != path:
            raise OSError("simulated atomic workbook replacement failure")
        return real_replace(source, target)

    monkeypatch.setattr(Path, "replace", fail_target_replace)
    with pytest.raises(OSError, match="simulated atomic workbook replacement failure"):
        migrate(path)

    assert path.read_bytes() == original
    assert list(tmp_path.glob("Mine-Services-Database-A3.*.xlsx")) == []
