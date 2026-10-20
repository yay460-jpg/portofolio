from openpyxl import load_workbook

from mine_services import schema
from mine_services.database_factory import ensure_a3_database
from mine_services.persistence import PersistenceStore


def test_fresh_install_creates_empty_valid_a3_workbook(tmp_path):
    target = tmp_path / "Database" / "Mine-Services-Database-A3.xlsx"

    result = ensure_a3_database(target)

    assert result == target
    assert target.is_file()
    workbook = load_workbook(target, data_only=True)
    try:
        assert workbook.sheetnames == schema.SHEETS
        assert workbook["_System"].cell(row=1, column=1).value == "schema_version"
        assert workbook["_System"].cell(row=1, column=2).value == "A.3"
        assert [cell.value for cell in workbook["NIK"][1]] == schema.NIK_HEADERS
        for entity in schema.DOMAIN_ENTITIES:
            assert [cell.value for cell in workbook[entity][1]] == schema.HEADERS[entity]
            assert workbook[entity].max_row == 1
        assert workbook["AuditLog"].max_row == 1
    finally:
        workbook.close()

    store = PersistenceStore(target)
    assert all(store.all(entity) == [] for entity in schema.DOMAIN_ENTITIES)
    assert store.audit() == []


def test_first_run_factory_never_overwrites_existing_database(tmp_path):
    target = tmp_path / "Database" / "Mine-Services-Database-A3.xlsx"
    target.parent.mkdir(parents=True)
    target.write_bytes(b"user-data-must-not-be-replaced")
    before = target.read_bytes()

    ensure_a3_database(target)

    assert target.read_bytes() == before
