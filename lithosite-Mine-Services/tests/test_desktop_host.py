import base64
import json
import sys
import pytest
from pathlib import Path
from urllib.request import Request, urlopen
from http.client import HTTPConnection
from urllib.error import HTTPError
from urllib.parse import urlencode

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))
sys.path.insert(0, str(Path(__file__).parents[1] / "desktop-host"))


def _load_server(monkeypatch, tmp_path):
    monkeypatch.setenv("MINE_SERVICES_DB", str(tmp_path / "desktop-host-a3.xlsx"))
    import server
    return server


def _make_test_database(path):
    from openpyxl import Workbook
    from mine_services.schema import DOMAIN_ENTITIES, HEADERS, SCHEMA_VERSION
    from mine_services.persistence import PersistenceStore

    wb = Workbook()
    wb.remove(wb.active)

    baseline = wb.create_sheet("_Baseline")
    baseline.append(["baseline_status", "LOCKED"])

    system = wb.create_sheet("_System")
    system.append(["schema_version", SCHEMA_VERSION])

    lists = wb.create_sheet("_Lists")
    defaults = PersistenceStore().controlled_lists
    headers = list(defaults)
    lists.append(headers)
    for i in range(max(len(values) for values in defaults.values())):
        lists.append([
            sorted(defaults[h])[i] if i < len(defaults[h]) else None
            for h in headers
        ])

    for entity in DOMAIN_ENTITIES:
        ws = wb.create_sheet(entity)
        ws.append(HEADERS[entity])

    audit = wb.create_sheet("AuditLog")
    audit.append(HEADERS["AuditLog"])
    wb.save(path)


def test_bounded_file_reader_enforces_exact_limit_and_sentinel(tmp_path):
    import server

    exact = tmp_path / "exact.bin"
    exact.write_bytes(b"12345")
    assert server._read_bounded_file(exact, 5) == b"12345"

    oversized = tmp_path / "oversized.bin"
    oversized.write_bytes(b"123456")
    with pytest.raises(OverflowError, match="File exceeds read limit"):
        server._read_bounded_file(oversized, 5)


def test_bounded_file_reader_rejects_negative_limit(tmp_path):
    import server

    sample = tmp_path / "sample.bin"
    sample.write_bytes(b"x")
    with pytest.raises(ValueError, match="Invalid file read limit"):
        server._read_bounded_file(sample, -1)


def test_json_body_reader_rejects_truncated_content(monkeypatch, tmp_path):
    import io
    import server

    assert server._read_exact_request_body(io.BytesIO(b'{"ok": true}'), 12) == b'{"ok": true}'
    with pytest.raises(ValueError, match="Incomplete request body"):
        server._read_exact_request_body(io.BytesIO(b'{"ok": true}'), 20)


def test_health_endpoint_is_offline_ready(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        host, port = httpd.server_address
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        with urlopen(f"http://{host}:{port}/health") as response:
            payload = json.loads(response.read().decode())
        assert payload["status"] == "READY"
        assert payload["offline"] is True
        assert payload["schema"] == "A.3"
        assert payload["runtime"] == "RuntimeAdapter"
        assert "database" not in payload
        assert str(db_path) not in json.dumps(payload)
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_runtime_endpoint_routes_to_adapter(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        request = Request(
            f"http://{host}:{port}/runtime",
            data=json.dumps({
                "request_id": "desktop-host-read",
                "operation": "READ",
                "entity": "Equipment",
            }).encode(),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urlopen(request) as response:
            payload = json.loads(response.read().decode())
        assert payload["request_id"] == "desktop-host-read"
        assert payload.get("status") in (None, "OK")
    finally:
        httpd.shutdown()
        httpd.server_close()



def test_runtime_rejects_invalid_and_over_limit_content_lengths_before_read(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address

        for length in ("0", "-1", "10000001", "not-a-number"):
            connection = HTTPConnection(host, port, timeout=3)
            try:
                connection.putrequest("POST", "/runtime")
                connection.putheader("Content-Length", length)
                connection.endheaders()
                response = connection.getresponse()
                payload = json.loads(response.read().decode("utf-8"))
                assert response.status == 400, (length, response.status, payload)
                assert payload["status"] == "REJECTED"
                assert payload["errors"][0]["code"] == "HOST-003"
                assert payload["errors"][0]["message"] == "Invalid JSON request"
            finally:
                connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_runtime_rejects_truncated_json_body_without_adapter_dispatch(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    def unexpected_dispatch(request):
        raise AssertionError("truncated body must not reach RuntimeAdapter")

    monkeypatch.setattr(server.ADAPTER, "handle", unexpected_dispatch)
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        connection = HTTPConnection(host, port, timeout=3)
        try:
            connection.putrequest("POST", "/runtime")
            connection.putheader("Content-Type", "application/json")
            connection.putheader("Content-Length", "100")
            connection.endheaders()
            connection.send(b'{"request_id":"truncated","operation":"READ"}')
            connection.sock.shutdown(1)
            response = connection.getresponse()
            payload = json.loads(response.read().decode("utf-8"))
            assert response.status == 400
            assert payload["status"] == "REJECTED"
            assert payload["errors"][0]["code"] == "HOST-003"
            assert payload["errors"][0]["message"] == "Invalid JSON request"
        finally:
            connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_user_guide_rejects_encoded_path_traversal(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    repo_root = tmp_path / "repo"
    guide_root = repo_root / "docs" / "lithosite" / "02_Mine-Services" / "10_User-Guides"
    guide_root.mkdir(parents=True)
    (tmp_path / "secret.pdf").write_bytes(b"private")
    monkeypatch.setattr(server, "REPO_ROOT", repo_root)

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        query = urlencode({"file": "../../../../../secret.pdf"})
        with pytest.raises(HTTPError) as rejected:
            urlopen(f"http://{host}:{port}/user-guide?{query}")
        assert rejected.value.code == 400
        payload = json.loads(rejected.value.read().decode("utf-8"))
        assert payload["status"] == "REJECTED"
        assert payload["errors"][0]["code"] == "HOST-007"
        assert "secret.pdf" not in json.dumps(payload)
    finally:
        httpd.shutdown()
        httpd.server_close()

@pytest.mark.parametrize("origin", ["http://evil.example", "http://localhost:8766", "http://127.0.0.1:5502", "null"])
def test_disallowed_origin_is_rejected(monkeypatch, tmp_path, origin):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        request = Request(
            f"http://{host}:{port}/runtime",
            data=json.dumps({"request_id":"origin-test","operation":"READ","entity":"Equipment"}).encode(),
            headers={"Content-Type":"application/json","Origin":origin},
            method="POST",
        )
        try:
            urlopen(request)
            assert False, "expected HTTP 403"
        except Exception as exc:
            assert "403" in str(exc)
    finally:
        httpd.shutdown()
        httpd.server_close()

def test_evidence_list_and_file_endpoints_serve_only_supported_record_files(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    evidence_root = tmp_path / "Evidence"
    monkeypatch.setattr(server, "EVIDENCE_ROOT", evidence_root)
    record_dir = evidence_root / "TargetPlan" / "PLN-TEST-001"
    record_dir.mkdir(parents=True)
    (record_dir / "mineout.pdf").write_bytes(b"%PDF-evidence-test")
    (record_dir / "pit-photo.PNG").write_bytes(b"\x89PNG\r\n evidence-test")
    (record_dir / "agreement.docx").write_bytes(b"word-document-test")
    (record_dir / "ignored.exe").write_bytes(b"not an evidence type")

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        base = f"http://{host}:{port}"
        query = urlencode({"module": "TargetPlan", "record_id": "PLN-TEST-001"})
        with urlopen(f"{base}/evidence/list?{query}") as response:
            listing = json.loads(response.read().decode("utf-8"))

        assert listing["status"] == "READY"
        assert listing["folder"] == "Database/Evidence/TargetPlan/PLN-TEST-001"
        assert {item["name"] for item in listing["files"]} == {
            "agreement.docx", "mineout.pdf", "pit-photo.PNG"
        }
        assert next(item for item in listing["files"] if item["name"] == "mineout.pdf")["previewable"] is True
        assert next(item for item in listing["files"] if item["name"] == "agreement.docx")["previewable"] is False

        with urlopen(f"{base}/evidence/file?{urlencode({'module':'TargetPlan','record_id':'PLN-TEST-001','filename':'mineout.pdf'})}") as response:
            pdf_bytes = response.read()
            assert response.headers.get_content_type() == "application/pdf"
            assert response.headers.get("Content-Disposition", "").startswith("inline;")
        assert pdf_bytes == b"%PDF-evidence-test"

        with urlopen(f"{base}/evidence/file?{urlencode({'module':'TargetPlan','record_id':'PLN-TEST-001','filename':'agreement.docx'})}") as response:
            docx_bytes = response.read()
            assert response.headers.get("Content-Disposition", "").startswith("attachment;")
        assert docx_bytes == b"word-document-test"
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_evidence_endpoints_reject_invalid_module_record_path_and_filename(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    monkeypatch.setattr(server, "EVIDENCE_ROOT", tmp_path / "Evidence")
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        base = f"http://{host}:{port}"
        invalid_queries = [
            ("/evidence/list", {"module": "Database", "record_id": "PLN-TEST-001"}),
            ("/evidence/list", {"module": "TargetPlan", "record_id": "../outside"}),
            ("/evidence/file", {"module": "TargetPlan", "record_id": "PLN-TEST-001", "filename": "../outside.pdf"}),
            ("/evidence/file", {"module": "TargetPlan", "record_id": "PLN-TEST-001", "filename": "payload.exe"}),
        ]
        for route, params in invalid_queries:
            try:
                urlopen(f"{base}{route}?{urlencode(params)}")
                assert False, f"expected request rejection for {route} {params}"
            except Exception as exc:
                assert "400" in str(exc)
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_evidence_endpoint_rejects_disallowed_origin(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    monkeypatch.setattr(server, "EVIDENCE_ROOT", tmp_path / "Evidence")
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        query = urlencode({"module": "TargetPlan", "record_id": "PLN-TEST-001"})
        request = Request(
            f"http://{host}:{port}/evidence/list?{query}",
            headers={"Origin": "http://evil.example"},
        )
        try:
            urlopen(request)
            assert False, "expected HTTP 403"
        except Exception as exc:
            assert "403" in str(exc)
    finally:
        httpd.shutdown()
        httpd.server_close()




def test_static_handler_serves_required_ui_assets_but_rejects_repository_files(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    repo_root = Path(__file__).parents[2].resolve()
    monkeypatch.setattr(server, "STATIC_ROOT", repo_root)
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        base = f"http://{host}:{port}"

        for allowed_path in (
            "/lithosite-Mine-Services/ui/shared/runtime-client.js",
            "/lithosite-Mine-Services/Artifacts/Mine-Services-Operations.html",
            "/lithosite-Mine-Services/shared/geo-engine.js",
            "/assets/lithosite-icons.svg",
            "/lithosite-Mine-Geologist/assets/icon-192.png",
        ):
            with urlopen(base + allowed_path) as response:
                assert response.status == 200, allowed_path
                assert response.read(), allowed_path

        for denied_path in (
            "/.gitignore",
            "/.github/workflows/mine-services-pytest.yml",
            "/lithosite-Mine-Services/src/mine_services/adapter.py",
            "/lithosite-Mine-Services/Documentation/Security/README.md",
            "/lithosite-Mine-Services/tests/test_runtime.py",
        ):
            with pytest.raises(HTTPError) as rejected:
                urlopen(base + denied_path)
            assert rejected.value.code == 404, denied_path
    finally:
        httpd.shutdown()
        httpd.server_close()


@pytest.mark.parametrize(
    ("module", "entity", "id_field"),
    [
        ("TargetPlan", "Plans", "plan_id"),
        ("HSE", "HSE", "hse_id"),
        ("Maintenance", "Maintenance", "maintenance_id"),
    ],
)
def test_evidence_owner_verification_uses_public_read_allowlist(
    monkeypatch, tmp_path, module, entity, id_field
):
    server = _load_server(monkeypatch, tmp_path)
    calls = []

    class FakeAdapter:
        def handle(self, request):
            calls.append(request)
            return {"request_id": request["request_id"], "status": "OK", "data": [{id_field: "record-1"}]}

    monkeypatch.setattr(server, "ADAPTER", FakeAdapter())

    assert server._evidence_record_exists(module, "record-1") is True
    assert len(calls) == 1
    assert calls[0]["operation"] == "READ"
    assert calls[0]["entity"] == entity


def test_evidence_owner_verification_rejects_unknown_module_before_read(monkeypatch, tmp_path):
    server = _load_server(monkeypatch, tmp_path)

    class NeverCalledAdapter:
        def handle(self, request):
            pytest.fail("unknown evidence module must not dispatch a READ")

    monkeypatch.setattr(server, "ADAPTER", NeverCalledAdapter())
    with pytest.raises(ValueError, match="Unsupported evidence module"):
        server._evidence_record_exists("AuditLog", "record-1")


@pytest.mark.parametrize(
    ("body", "internal_detail"),
    [
        (json.dumps([]).encode("utf-8"), "Report model must be an object"),
        (b'{"broken":', "Expecting value"),
    ],
)
def test_report_pdf_rejects_malformed_requests_with_sanitized_error(
    monkeypatch, tmp_path, body, internal_detail
):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        request = Request(
            f"http://{host}:{port}/report-pdf",
            data=body,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with pytest.raises(HTTPError) as rejected:
            urlopen(request)
        assert rejected.value.code == 400
        payload = json.loads(rejected.value.read().decode("utf-8"))
        assert payload["status"] == "REJECTED"
        assert payload["errors"][0]["code"] == "HOST-008"
        assert payload["errors"][0]["message"] == "Invalid report request"
        assert internal_detail not in json.dumps(payload)
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_evidence_rejects_symlinked_module_and_file_paths(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    evidence_root = tmp_path / "Evidence"
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "secret.pdf").write_bytes(b"%PDF-private")
    evidence_root.mkdir()
    try:
        (evidence_root / "TargetPlan").symlink_to(outside, target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("directory symlinks are not available on this platform")

    monkeypatch.setattr(server, "EVIDENCE_ROOT", evidence_root)
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        base = f"http://{host}:{port}"

        with pytest.raises(HTTPError) as module_rejected:
            urlopen(
                f"{base}/evidence/list?"
                + urlencode({"module": "TargetPlan", "record_id": "PLN-TEST-001"})
            )
        assert module_rejected.value.code == 400
        module_payload = json.loads(module_rejected.value.read().decode("utf-8"))
        assert module_payload["status"] == "REJECTED"
        assert "secret.pdf" not in json.dumps(module_payload)

        # Replace the module link with an in-root record and a file link that
        # points outside Evidence; the file endpoint must refuse to serve it.
        (evidence_root / "TargetPlan").unlink()
        record_dir = evidence_root / "TargetPlan" / "PLN-TEST-001"
        record_dir.mkdir(parents=True)
        try:
            (record_dir / "linked.pdf").symlink_to(outside / "secret.pdf")
        except (OSError, NotImplementedError):
            pytest.skip("file symlinks are not available on this platform")

        with pytest.raises(HTTPError) as file_rejected:
            urlopen(
                f"{base}/evidence/file?"
                + urlencode({
                    "module": "TargetPlan",
                    "record_id": "PLN-TEST-001",
                    "filename": "linked.pdf",
                })
            )
        assert file_rejected.value.code == 404
        file_payload = json.loads(file_rejected.value.read().decode("utf-8"))
        assert file_payload["status"] == "REJECTED"
        assert "secret.pdf" not in json.dumps(file_payload)
    finally:
        httpd.shutdown()
        httpd.server_close()


@pytest.mark.parametrize(
    "body",
    [
        b"{malformed-json",
        b"[]",
        b"null",
    ],
)
def test_report_pdf_rejects_malformed_or_non_object_request_without_details(
    monkeypatch, tmp_path, body
):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        connection = HTTPConnection(host, port, timeout=3)
        try:
            connection.putrequest("POST", "/report-pdf")
            connection.putheader("Content-Type", "application/json")
            connection.putheader("Content-Length", str(len(body)))
            connection.endheaders()
            connection.send(body)
            response = connection.getresponse()
            response_body = response.read().decode("utf-8")
            payload = json.loads(response_body)
            assert response.status == 400
            assert payload["status"] == "REJECTED"
            assert payload["errors"][0]["code"] == "HOST-008"
            assert payload["errors"][0]["message"] == "Invalid report request"
            assert "Traceback" not in response_body
            assert str(db_path) not in response_body
        finally:
            connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()


@pytest.mark.parametrize("length", ["-1", "10000001", "not-a-number"])
def test_report_pdf_rejects_invalid_or_over_limit_content_length_before_read(
    monkeypatch, tmp_path, length
):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        connection = HTTPConnection(host, port, timeout=3)
        try:
            connection.putrequest("POST", "/report-pdf")
            connection.putheader("Content-Type", "application/json")
            connection.putheader("Content-Length", length)
            connection.endheaders()
            response = connection.getresponse()
            response_body = response.read().decode("utf-8")
            payload = json.loads(response_body)
            assert response.status == 400
            assert payload["status"] == "REJECTED"
            assert payload["errors"][0]["code"] == "HOST-008"
            assert payload["errors"][0]["message"] == "Invalid report request"
            assert "Traceback" not in response_body
            assert str(db_path) not in response_body
        finally:
            connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()


@pytest.mark.parametrize(
    "model",
    [
        {"period": ["unexpected-list"]},
        {"kpi": ["unexpected-list"]},
        {"source_counts": ["unexpected-list"]},
        {"section_data": ["unexpected-list"], "sections": ["Executive Summary"]},
    ],
)
def test_report_pdf_rejects_wrongly_typed_nested_payload_without_details(
    monkeypatch, tmp_path, model
):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    body = json.dumps(model).encode("utf-8")
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        connection = HTTPConnection(host, port, timeout=3)
        try:
            connection.putrequest("POST", "/report-pdf")
            connection.putheader("Content-Type", "application/json")
            connection.putheader("Content-Length", str(len(body)))
            connection.endheaders()
            connection.send(body)
            response = connection.getresponse()
            response_body = response.read().decode("utf-8")
            payload = json.loads(response_body)
            assert response.status == 400
            assert payload["status"] == "REJECTED"
            assert payload["errors"][0]["code"] == "HOST-008"
            assert payload["errors"][0]["message"] == "Invalid report request"
            assert "Traceback" not in response_body
            assert str(db_path) not in response_body
        finally:
            connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()



def test_report_pdf_rejects_model_over_processing_limit(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    body = json.dumps({"large": "x" * 1_000_001}).encode("utf-8")
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        request = Request(f"http://{host}:{port}/report-pdf", data=body,
                          headers={"Content-Type": "application/json"}, method="POST")
        with pytest.raises(HTTPError) as rejected:
            urlopen(request, timeout=5)
        assert rejected.value.code == 400
        response = rejected.value.read().decode("utf-8")
        payload = json.loads(response)
        assert payload["status"] == "REJECTED"
        assert payload["errors"][0]["code"] == "HOST-008"
        assert payload["errors"][0]["message"] == "Invalid report request"
        assert "Traceback" not in response
        assert str(db_path) not in response
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_report_pdf_accepts_small_valid_object_model(monkeypatch, tmp_path):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    body = json.dumps({"title": "Operations Report"}).encode("utf-8")
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        request = Request(f"http://{host}:{port}/report-pdf", data=body,
                          headers={"Content-Type": "application/json"}, method="POST")
        with urlopen(request, timeout=5) as response:
            payload = json.loads(response.read().decode("utf-8"))
        assert payload["status"] == "READY"
        assert payload["mime"] == "application/pdf"
        assert base64.b64decode(payload["data"]).startswith(b"%PDF-")
    finally:
        httpd.shutdown()
        httpd.server_close()

@pytest.mark.parametrize("method", ["PUT", "PATCH", "DELETE"])
def test_unsupported_http_methods_do_not_dispatch_runtime_operations(
    monkeypatch, tmp_path, method
):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        connection = HTTPConnection(host, port, timeout=3)
        try:
            connection.request(method, "/runtime")
            response = connection.getresponse()
            response_body = response.read().decode("utf-8", errors="replace")
            assert response.status == 501
            assert "Traceback" not in response_body
            assert str(db_path) not in response_body
        finally:
            connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()


@pytest.mark.parametrize("path", ["/health", "/assets/lithosite-icons.svg"])
def test_untrusted_origin_gets_no_cors_read_permission_on_public_get_routes(
    monkeypatch, tmp_path, path
):
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    repo_root = Path(__file__).parents[2].resolve()
    monkeypatch.setattr(server, "STATIC_ROOT", repo_root)
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        request = Request(
            f"http://{host}:{port}{path}",
            headers={"Origin": "http://evil.example"},
        )
        with urlopen(request, timeout=3) as response:
            body = response.read()
            assert response.status == 200
            assert response.headers.get("Access-Control-Allow-Origin") is None
            if path == "/health":
                payload = json.loads(body.decode("utf-8"))
                assert set(payload) == {"status", "offline", "schema", "runtime"}
            else:
                assert body
    finally:
        httpd.shutdown()
        httpd.server_close()


def test_untrusted_origin_cannot_read_any_remaining_get_route(monkeypatch, tmp_path):
    """Exercise GET Origin/CORS behavior across route-specific handler branches."""
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    # Use a harmless fixture for the successful user-guide path.
    guide_root = tmp_path / "docs" / "lithosite" / "02_Mine-Services" / "10_User-Guides"
    guide_root.mkdir(parents=True)
    (guide_root / "sample.pdf").write_bytes(b"%PDF-test")
    monkeypatch.setattr(server, "REPO_ROOT", tmp_path)
    monkeypatch.setattr(server, "STATIC_ROOT", Path(__file__).parents[2].resolve())
    monkeypatch.setattr(server, "EVIDENCE_ROOT", tmp_path / "Evidence")

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        cases = [
            ("/", 302),
            ("/health", 200),
            ("/user-guide?file=sample.pdf", 200),
            ("/nik/profile?nik_id=NIK-TEST-001", 403),
            ("/evidence/status?module=TargetPlan", 403),
            ("/evidence/list?module=TargetPlan&record_id=PLN-TEST-001", 403),
            ("/evidence/file?module=TargetPlan&record_id=PLN-TEST-001&filename=proof.pdf", 403),
            ("/assets/lithosite-icons.svg", 200),
        ]
        # Use a fresh connection per route because the host uses HTTP/1.0
        # connection-close semantics for several response paths.
        for route, expected_status in cases:
            connection = HTTPConnection(host, port, timeout=3)
            try:
                connection.request(
                    "GET",
                    route,
                    headers={"Origin": "http://evil.example"},
                )
                response = connection.getresponse()
                body = response.read()
                assert response.status == expected_status, route
                assert response.getheader("Access-Control-Allow-Origin") is None, route
                if expected_status == 200:
                    assert body, route
            finally:
                connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()


@pytest.mark.parametrize(
    ("route", "body"),
    [
        ("/nik/clear?nik_id=NIK-TEST-001", None),
        ("/nik/upload?nik_id=NIK-TEST-001", b"not-an-image"),
        ("/evidence/upload?module=TargetPlan&record_id=PLN-TEST-001&filename=proof.pdf", b"proof"),
        ("/evidence/preview", b"{}"),
        ("/report-pdf", b"{}"),
    ],
)
def test_untrusted_origin_is_rejected_before_post_route_processing(
    monkeypatch, tmp_path, route, body
):
    """Every dedicated POST route must reject untrusted Origin before side effects."""
    db_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(db_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(db_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")
    import server

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    try:
        import threading
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        host, port = httpd.server_address
        connection = HTTPConnection(host, port, timeout=3)
        headers = {"Origin": "http://evil.example"}
        if body is not None:
            headers["Content-Type"] = "application/json" if route == "/evidence/preview" or route == "/report-pdf" else "image/jpeg"
        try:
            connection.request("POST", route, body=body, headers=headers)
            response = connection.getresponse()
            response_body = response.read().decode("utf-8", errors="replace")
            assert response.status == 403, route
            assert response.getheader("Access-Control-Allow-Origin") is None, route
            assert "Traceback" not in response_body, route
            assert str(db_path) not in response_body, route
        finally:
            connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()

@pytest.mark.parametrize(
    ("origin", "expected_allow_origin"),
    [
        ("http://127.0.0.1:8765", "http://127.0.0.1:8765"),
        ("http://localhost:8765", "http://localhost:8765"),
        ("http://evil.example", None),
        ("null", None),
        (None, None),
    ],
)
def test_options_preflight_only_grants_allowlisted_origin(
    monkeypatch, tmp_path, origin, expected_allow_origin
):
    """OPTIONS advertises CORS read permission only to allowlisted browser origins."""
    import threading
    import server

    monkeypatch.setenv("MINE_SERVICES_PORT", "8765")
    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    host, port = httpd.server_address
    try:
        connection = HTTPConnection(host, port, timeout=5)
        headers = {"Access-Control-Request-Method": "POST"}
        if origin is not None:
            headers["Origin"] = origin
        connection.request("OPTIONS", "/runtime", headers=headers)
        response = connection.getresponse()
        response.read()
        assert response.status == 204
        assert response.getheader("Access-Control-Allow-Origin") == expected_allow_origin
        if expected_allow_origin is None:
            assert response.getheader("Vary") is None
        else:
            assert response.getheader("Vary") == "Origin"
        assert response.getheader("Access-Control-Allow-Methods") == "GET, POST, OPTIONS"
        assert response.getheader("Access-Control-Allow-Headers") == "Content-Type"
        connection.close()
    finally:
        httpd.shutdown()
        httpd.server_close()
        thread.join(timeout=3)

