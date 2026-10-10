import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))

from mine_services import (
    ApplicationService,
    AuditRepository,
    ImportCoordinator,
    PersistenceStore,
    SnapshotManager,
)


def valid_equipment(equipment_id="EQ-01"):
    return {
        "equipment_id": equipment_id,
        "unit_no": "DT-001",
        "category": "Heavy Equipment",
        "type": "Dump Truck",
        "owner_type": "Owner",
        "status": "Active",
    }


def valid_work_front(work_front_id="WF-01"):
    return {
        "work_front_id": work_front_id,
        "domain": "Road & Hauling",
        "location": "Pit A",
        "status": "Active",
    }


def seed():
    app = ApplicationService()
    assert app.create("Equipment", valid_equipment(), "s1")["status"] == "COMMITTED"
    assert app.create("WorkFront", valid_work_front(), "s2")["status"] == "COMMITTED"
    return app


def test_update_read_and_delete():
    app = seed()
    assert app.update("Equipment", "EQ-01", {"status": "Inactive"}, "u1")["status"] == "COMMITTED"
    assert app.read("Equipment", "EQ-01")["status"] == "Inactive"
    assert app.delete("Equipment", "EQ-01", "d1")["status"] == "COMMITTED"
    assert app.read("Equipment", "EQ-01") is None
    assert app.store.audit()[-1]["action"] == "DELETE"


def test_import_is_atomic_on_validation_error():
    store = PersistenceStore()
    imp = ImportCoordinator(store)
    result = imp.import_dataset(
        {
            "Equipment": [
                valid_equipment("EQ-01"),
                valid_equipment("EQ-01"),
            ]
        }
    )
    assert result["status"] == "REJECTED"
    assert store.all("Equipment") == []


def test_import_commits_valid_dataset():
    store = PersistenceStore()
    imp = ImportCoordinator(store)
    result = imp.import_dataset({"Equipment": [valid_equipment("EQ-01")]})
    assert result["status"] == "COMMITTED"
    assert store.exists("Equipment", "EQ-01")


def test_snapshot_checksum_and_restore():
    app = seed()
    snap = SnapshotManager(app.store).capture()
    assert SnapshotManager(app.store).verify(snap)

    assert app.create("Equipment", valid_equipment("EQ-02"), "s3")["status"] == "COMMITTED"
    assert SnapshotManager(app.store).restore(snap)["status"] == "COMMITTED"
    assert app.read("Equipment", "EQ-02") is None


def test_snapshot_tamper_rejected():
    app = seed()
    snap = SnapshotManager(app.store).capture()
    snap["payload"]["data"]["Equipment"]["EQ-01"]["type"] = "Tampered"
    assert SnapshotManager(app.store).restore(snap)["status"] == "REJECTED"


def test_import_allows_parent_child_cross_fk_in_same_dataset():
    store = PersistenceStore()
    result = ImportCoordinator(store).import_dataset({
        "GlobalCapacity": [{
            "capacity_profile_id": "GC-CROSS",
            "capacity_name": "Cross-FK Test Capacity",
            "capacity_value": 25.0,
            "measurement": "ton",
            "status": "Active",
        }],
        "Equipment": [{
            **valid_equipment("EQ-CROSS"),
            "capacity_profile_id": "GC-CROSS",
        }],
        "WorkFront": [valid_work_front("WF-CROSS")],
        "Maintenance": [{
            "maintenance_id": "M-CROSS",
            "equipment_id": "EQ-CROSS",
            "event_type": "Preventive",
            "status": "Open",
        }],
        "Operations": [{
            "transaction_id": "TX-CROSS",
            "domain": "Road & Hauling",
            "work_front_id": "WF-CROSS",
            "equipment_id": "EQ-CROSS",
            "activity": "Hauling",
            "retase": 1,
            "capacity_profile_id": "GC-CROSS",
            "applied_capacity": 25.0,
            "capacity_measurement": "ton",
            "quantity": 25.0,
            "measurement": "ton",
            "status": "DRAFT",
        }],
    })
    assert result["status"] == "COMMITTED"
    assert store.exists("Maintenance", "M-CROSS")
    assert store.exists("Operations", "TX-CROSS")


def test_delete_reference_is_rejected():
    app = seed()
    assert app.create(
        "Maintenance",
        {
            "maintenance_id": "M-REF",
            "equipment_id": "EQ-01",
            "event_type": "Preventive",
            "status": "Open",
        },
        "create-ref",
    )["status"] == "COMMITTED"
    result = app.delete("Equipment", "EQ-01", "delete-ref")
    assert result["status"] == "REJECTED"
    assert result["errors"][0]["code"] == "APP-004"


def test_audit_failure_rolls_back_mutation():
    app = ApplicationService()
    original = app.audit_repository.append

    def fail_audit(**kwargs):
        raise RuntimeError("AUDIT_FAILURE")

    app.audit_repository.append = fail_audit
    try:
        result = app.create("Equipment", valid_equipment("EQ-AUDIT"), "audit-fail")
    finally:
        app.audit_repository.append = original

    assert result["status"] == "REJECTED"
    assert result["errors"][0]["code"] == "APP-005"
    assert app.read("Equipment", "EQ-AUDIT") is None


def test_application_service_uses_audit_repository_boundary():
    store = PersistenceStore()
    audit = AuditRepository(store)
    app = ApplicationService(store=store, audit_repository=audit)

    result = app.create("Equipment", valid_equipment("EQ-AUDIT-BOUNDARY"), "audit-boundary")

    assert result["status"] == "COMMITTED"
    assert audit.all()[-1]["action"] == "CREATE"
    assert audit.all()[-1]["source"] == "audit-boundary"


def test_request_id_is_required_and_idempotent():
    app = ApplicationService()
    missing = app.create("Equipment", valid_equipment("EQ-REQ"), "")
    assert missing["status"] == "REJECTED"
    assert missing["errors"][0]["code"] == "APP-001"

    first = app.create("Equipment", valid_equipment("EQ-IDEMP"), "same-request")
    second = app.create("Equipment", valid_equipment("EQ-IDEMP-2"), "same-request")
    assert first["status"] == "COMMITTED"
    assert second["status"] == "DUPLICATE_REQUEST"


def test_application_not_found_returns_app_error():
    app = ApplicationService()
    result = app.update("Equipment", "NO-SUCH", {"status": "Inactive"}, "missing")
    assert result["status"] == "REJECTED"
    assert result["errors"][0]["code"] == "APP-002"


def test_snapshot_manifest_and_restore_audit():
    app = seed()
    manager = SnapshotManager(app.store)
    snap = manager.capture(source="Stage-3-Test")
    assert snap["schema_version"] == "A.3"
    assert snap["source"] == "Stage-3-Test"
    assert snap["audit_included"] is True
    assert snap["entity_counts"]["Equipment"] == 1
    assert snap["status"] == "SEALED"

    result = manager.restore(snap)
    assert result["status"] == "COMMITTED"
    assert app.store.audit()[-1]["action"] == "RESTORE"


def test_snapshot_dry_run_does_not_mutate():
    app = seed()
    manager = SnapshotManager(app.store)
    snap = manager.capture()
    assert app.create("Equipment", valid_equipment("EQ-DRY"), "dry-seed")["status"] == "COMMITTED"
    before = app.read("Equipment", "EQ-DRY")
    result = manager.restore(snap, mode="DRY_RUN")
    assert result["status"] == "VALIDATED"
    assert app.read("Equipment", "EQ-DRY") == before


def test_snapshot_entity_counts_tamper_is_rejected():
    app = seed()
    manager = SnapshotManager(app.store)
    snap = manager.capture()
    snap["entity_counts"]["Equipment"] = 999
    assert manager.restore(snap)["status"] == "REJECTED"


def test_snapshot_merge_rejects_primary_key_collision():
    app = seed()
    manager = SnapshotManager(app.store)
    snap = manager.capture()
    result = manager.restore(snap, mode="MERGE_RUNTIME")
    assert result["status"] == "REJECTED"
    assert result["reason"] == "MERGE_PK_COLLISION"


def test_xlsx_persistence_roundtrip(tmp_path):
    from mine_services.schema import DOMAIN_ENTITIES, HEADERS, SCHEMA_VERSION

    path = tmp_path / "runtime.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    baseline = wb.create_sheet("_Baseline")
    baseline.append(["baseline_status", "LOCKED"])

    system = wb.create_sheet("_System")
    system.append(["schema_version", SCHEMA_VERSION])

    lists = wb.create_sheet("_Lists")
    defaults = PersistenceStore().controlled_lists
    list_headers = list(defaults)
    lists.append(list_headers)
    for row_index in range(max(len(values) for values in defaults.values())):
        lists.append([
            sorted(defaults[header])[row_index] if row_index < len(defaults[header]) else None
            for header in list_headers
        ])

    for entity in DOMAIN_ENTITIES:
        ws = wb.create_sheet(entity)
        ws.append(HEADERS[entity])

    audit = wb.create_sheet("AuditLog")
    audit.append(HEADERS["AuditLog"])
    wb.save(path)

    store = PersistenceStore(path)
    result = ApplicationService(store).create(
        "Equipment",
        valid_equipment("EQ-XLSX"),
        "xlsx-roundtrip",
    )
    assert result["status"] == "COMMITTED"

    reloaded = PersistenceStore(path)
    assert reloaded.exists("Equipment", "EQ-XLSX")
    assert reloaded.get("Equipment", "EQ-XLSX")["equipment_id"] == "EQ-XLSX"

    verify = __import__("openpyxl").load_workbook(path, data_only=True)
    assert verify["_Baseline"]["B1"].value == "LOCKED"

