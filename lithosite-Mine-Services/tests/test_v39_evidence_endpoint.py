import base64
import importlib.util
import json
import sys
import threading
import uuid
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

import pytest

ROOT = Path(__file__).resolve().parents[1]
HOST_FILE = ROOT / "desktop-host" / "server.py"
sys.path.insert(0, str(ROOT / "src"))


def _make_test_database(path):
    from openpyxl import Workbook
    from mine_services.persistence import PersistenceStore
    from mine_services.schema import DOMAIN_ENTITIES, HEADERS, SCHEMA_VERSION

    workbook = Workbook()
    workbook.remove(workbook.active)

    baseline = workbook.create_sheet("_Baseline")
    baseline.append(["baseline_status", "LOCKED"])

    system = workbook.create_sheet("_System")
    system.append(["schema_version", SCHEMA_VERSION])

    lists = workbook.create_sheet("_Lists")
    defaults = PersistenceStore().controlled_lists
    headers = list(defaults)
    lists.append(headers)
    for index in range(max(len(values) for values in defaults.values())):
        lists.append([
            sorted(defaults[header])[index] if index < len(defaults[header]) else None
            for header in headers
        ])

    for entity in DOMAIN_ENTITIES:
        worksheet = workbook.create_sheet(entity)
        worksheet.append(HEADERS[entity])

    audit = workbook.create_sheet("AuditLog")
    audit.append(HEADERS["AuditLog"])
    workbook.save(path)


@pytest.fixture
def evidence_host(monkeypatch, tmp_path):
    database_path = tmp_path / "desktop-host-a3.xlsx"
    _make_test_database(database_path)
    monkeypatch.setenv("MINE_SERVICES_DB", str(database_path))
    monkeypatch.setenv("MINE_SERVICES_SCHEMA", "A3")

    module_name = f"_evidence_test_server_{uuid.uuid4().hex}"
    spec = importlib.util.spec_from_file_location(module_name, HOST_FILE)
    assert spec is not None and spec.loader is not None
    server = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = server
    spec.loader.exec_module(server)

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    host, port = httpd.server_address

    try:
        yield f"http://{host}:{port}", server
    finally:
        httpd.shutdown()
        httpd.server_close()
        thread.join(timeout=3)
        sys.modules.pop(module_name, None)


def _get(base_url, endpoint, params):
    separator = "&" if "?" in endpoint else "?"
    url = f"{base_url}{endpoint}{separator}{urlencode(params)}"
    try:
        with urlopen(url, timeout=5) as response:
            return response.status, response.headers, response.read()
    except HTTPError as error:
        return error.code, error.headers, error.read()


def _payload(body):
    return json.loads(body.decode("utf-8"))


def _post_json(base_url, endpoint, payload):
    request = Request(
        f"{base_url}{endpoint}",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urlopen(request, timeout=5) as response:
            return response.status, response.headers, response.read()
    except HTTPError as error:
        return error.code, error.headers, error.read()


def _post_raw(base_url, endpoint, params, body, content_type="application/octet-stream"):
    url = f"{base_url}{endpoint}?{urlencode(params)}"
    request = Request(
        url,
        data=body,
        headers={"Content-Type": content_type},
        method="POST",
    )
    try:
        with urlopen(request, timeout=10) as response:
            return response.status, response.headers, response.read()
    except HTTPError as error:
        return error.code, error.headers, error.read()


def test_evidence_upload_saves_to_the_plan_folder_and_never_overwrites(evidence_host, monkeypatch):
    base_url, server = evidence_host
    monkeypatch.setattr(server, "_evidence_record_exists", lambda module, record_id: module == "TargetPlan" and record_id == "PLN-UPLOAD-01")
    params = {
        "module": "TargetPlan",
        "record_id": "PLN-UPLOAD-01",
        "filename": "Mine-out agreement.pdf",
    }
    original = b"%PDF-1.7 accepted evidence"
    status, headers, body = _post_raw(base_url, "/evidence/upload", params, original)

    assert status == 201
    assert headers.get_content_type() == "application/json"
    payload = _payload(body)
    assert payload["status"] == "READY"
    assert payload["folder"] == "Database/Evidence/TargetPlan/PLN-UPLOAD-01"
    assert payload["filename"] == params["filename"]
    saved = server.EVIDENCE_ROOT / "TargetPlan" / "PLN-UPLOAD-01" / params["filename"]
    assert saved.read_bytes() == original

    duplicate_status, _duplicate_headers, duplicate_body = _post_raw(
        base_url, "/evidence/upload", params, b"%PDF-1.7 replacement"
    )
    assert duplicate_status == 409
    assert _payload(duplicate_body)["status"] == "REJECTED"
    assert saved.read_bytes() == original


@pytest.mark.parametrize(
    ("params", "record_exists", "expected_status"),
    [
        ({"module": "Unknown", "record_id": "REC-01", "filename": "evidence.jpg"}, True, 400),
        ({"module": "TargetPlan", "record_id": "../outside", "filename": "evidence.pdf"}, True, 400),
        ({"module": "TargetPlan", "record_id": "PLN-UPLOAD-02", "filename": "../outside.pdf"}, True, 400),
        ({"module": "TargetPlan", "record_id": "PLN-UPLOAD-02", "filename": "payload.exe"}, True, 400),
        ({"module": "TargetPlan", "record_id": "PLN-UPLOAD-02", "filename": "evidence.pdf"}, False, 404),
        ({"module": "TargetPlan", "record_id": "PLN-UPLOAD-02", "filename": "empty.pdf"}, True, 400),
    ],
)
def test_evidence_upload_rejects_invalid_module_record_filename_or_empty_body(
    evidence_host, monkeypatch, params, record_exists, expected_status
):
    base_url, server = evidence_host
    monkeypatch.setattr(server, "_evidence_record_exists", lambda _module, _record_id: record_exists)
    body = b"" if params["filename"] == "empty.pdf" else b"sample bytes"

    status, _headers, response_body = _post_raw(base_url, "/evidence/upload", params, body)

    assert status == expected_status
    result = _payload(response_body)
    assert result["status"] == "REJECTED"
    assert result["errors"]


@pytest.mark.parametrize(
    ("module", "record_id"),
    [
        ("HSE", "HSE-UPLOAD-01"),
        ("Maintenance", "MAINT-UPLOAD-01"),
    ],
)
def test_evidence_upload_accepts_hse_and_maintenance_record_scoped_files(
    evidence_host, monkeypatch, module, record_id
):
    base_url, server = evidence_host
    monkeypatch.setattr(
        server,
        "_evidence_record_exists",
        lambda candidate_module, candidate_id: (candidate_module, candidate_id) == (module, record_id),
    )
    params = {"module": module, "record_id": record_id, "filename": "evidence.jpg"}
    original = b"module-specific evidence"

    status, _headers, body = _post_raw(base_url, "/evidence/upload", params, original)

    assert status == 201
    payload = _payload(body)
    assert payload["status"] == "READY"
    assert payload["folder"] == f"Database/Evidence/{module}/{record_id}"
    saved = server.EVIDENCE_ROOT / module / record_id / params["filename"]
    assert saved.read_bytes() == original



def test_evidence_status_lists_only_records_with_supported_direct_child_files(evidence_host):
    base_url, server = evidence_host
    fixtures = {
        "TargetPlan": ("PLN-STATUS-01", ["approval.pdf", "field-photo.JPG"], ["ignored.exe"]),
        "HSE": ("HSE-STATUS-01", ["incident.jpg"], ["notes.tmp"]),
        "Maintenance": ("MNT-STATUS-01", ["repair.docx"], ["backup.zip"]),
    }
    for module, (record_id, valid_names, ignored_names) in fixtures.items():
        record_dir = server.EVIDENCE_ROOT / module / record_id
        record_dir.mkdir(parents=True)
        for name in valid_names + ignored_names:
            (record_dir / name).write_bytes(b"test evidence")
        (record_dir / "subfolder").mkdir()
        (record_dir / "subfolder" / "nested.pdf").write_bytes(b"nested file")
        (server.EVIDENCE_ROOT / module / (record_id + "-EMPTY")).mkdir()

        status, headers, body = _get(
            base_url,
            "/evidence/status",
            {"module": module},
        )

        assert status == 200
        assert headers.get_content_type() == "application/json"
        payload = _payload(body)
        assert payload["status"] == "READY"
        assert payload["module"] == module
        assert payload["records"] == [{"record_id": record_id, "file_count": len(valid_names)}]


def test_evidence_status_rejects_unsupported_modules(evidence_host):
    base_url, _server = evidence_host

    status, headers, body = _get(
        base_url,
        "/evidence/status",
        {"module": "Unknown"},
    )

    assert status == 400
    assert headers.get_content_type() == "application/json"
    payload = _payload(body)
    assert payload["status"] == "REJECTED"
    assert payload["errors"][0]["code"] == "HOST-009"


def test_listing_a_missing_evidence_folder_does_not_create_it(evidence_host):
    base_url, server = evidence_host
    missing = server.EVIDENCE_ROOT / "TargetPlan" / "PLN-NO-FOLDER-01"
    assert not missing.exists()

    status, _headers, body = _get(
        base_url,
        "/evidence/list",
        {"module": "TargetPlan", "record_id": "PLN-NO-FOLDER-01"},
    )

    assert status == 200
    assert _payload(body)["files"] == []
    assert not missing.exists()


def test_runtime_plan_create_and_delete_manage_only_its_evidence_folder(evidence_host, monkeypatch):
    base_url, server = evidence_host
    evidence_root = server.EVIDENCE_ROOT
    plan_id = "PLN-LIFECYCLE-01"
    plan_folder = evidence_root / "TargetPlan" / plan_id
    sibling_folder = evidence_root / "TargetPlan" / "PLN-KEEP-01"

    class CommittingAdapter:
        def handle(self, request):
            if request.get("operation") == "DELETE":
                assert plan_folder.is_dir(), "Evidence must remain until RuntimeAdapter commits deletion"
            return {"request_id": request.get("request_id"), "status": "COMMITTED"}

    monkeypatch.setattr(server, "ADAPTER", CommittingAdapter())
    status, _headers, body = _post_json(
        base_url,
        "/runtime",
        {
            "request_id": "test-plan-create",
            "operation": "CREATE",
            "entity": "Plans",
            "row": {"plan_id": plan_id},
        },
    )
    assert status == 200
    created = _payload(body)
    assert created["status"] == "COMMITTED"
    assert created["evidence_folder_status"] == "READY"
    assert plan_folder.is_dir()

    (plan_folder / "approval.pdf").write_bytes(b"evidence")
    sibling_folder.mkdir(parents=True)
    (sibling_folder / "keep.pdf").write_bytes(b"keep")

    status, _headers, body = _post_json(
        base_url,
        "/runtime",
        {
            "request_id": "test-plan-delete",
            "operation": "DELETE",
            "entity": "Plans",
            "entity_id": plan_id,
        },
    )
    assert status == 200
    deleted = _payload(body)
    assert deleted["status"] == "COMMITTED"
    assert deleted["evidence_cleanup_status"] == "CLEANED"
    assert not plan_folder.exists()
    assert (sibling_folder / "keep.pdf").read_bytes() == b"keep"



@pytest.mark.parametrize(
    ("entity", "module", "record_id", "id_field"),
    [
        ("HSE", "HSE", "HSE-LIFECYCLE-01", "hse_id"),
        ("Maintenance", "Maintenance", "MAINT-LIFECYCLE-01", "maintenance_id"),
    ],
)
def test_runtime_hse_and_maintenance_mutations_manage_only_their_evidence_folder(
    evidence_host, monkeypatch, entity, module, record_id, id_field
):
    base_url, server = evidence_host
    record_folder = server.EVIDENCE_ROOT / module / record_id
    sibling_folder = server.EVIDENCE_ROOT / module / (record_id + "-KEEP")

    class CommittingAdapter:
        def handle(self, request):
            if request.get("operation") == "DELETE":
                assert record_folder.is_dir(), "Evidence must remain until deletion is committed"
            return {"request_id": request.get("request_id"), "status": "COMMITTED"}

    monkeypatch.setattr(server, "ADAPTER", CommittingAdapter())
    status, _headers, body = _post_json(
        base_url,
        "/runtime",
        {
            "request_id": f"test-{module.casefold()}-create",
            "operation": "CREATE",
            "entity": entity,
            "row": {id_field: record_id},
        },
    )
    assert status == 200
    created = _payload(body)
    assert created["status"] == "COMMITTED"
    assert created["evidence_folder_status"] == "READY"
    assert record_folder.is_dir()

    (record_folder / "evidence.pdf").write_bytes(b"record evidence")
    sibling_folder.mkdir(parents=True)
    (sibling_folder / "keep.pdf").write_bytes(b"keep")

    status, _headers, body = _post_json(
        base_url,
        "/runtime",
        {
            "request_id": f"test-{module.casefold()}-delete",
            "operation": "DELETE",
            "entity": entity,
            "entity_id": record_id,
        },
    )
    assert status == 200
    deleted = _payload(body)
    assert deleted["status"] == "COMMITTED"
    assert deleted["evidence_cleanup_status"] == "CLEANED"
    assert not record_folder.exists()
    assert (sibling_folder / "keep.pdf").read_bytes() == b"keep"


def test_runtime_rejected_plan_delete_keeps_evidence_folder(evidence_host, monkeypatch):
    base_url, server = evidence_host
    plan_id = "PLN-DELETE-REJECTED"
    plan_folder = server.EVIDENCE_ROOT / "TargetPlan" / plan_id
    plan_folder.mkdir(parents=True)
    (plan_folder / "approval.pdf").write_bytes(b"keep evidence")

    class RejectingAdapter:
        def handle(self, request):
            return {
                "request_id": request.get("request_id"),
                "status": "REJECTED",
                "errors": [{"code": "VAL-E005", "message": "Plan is still referenced"}],
            }

    monkeypatch.setattr(server, "ADAPTER", RejectingAdapter())
    status, _headers, body = _post_json(
        base_url,
        "/runtime",
        {
            "request_id": "test-plan-delete-rejected",
            "operation": "DELETE",
            "entity": "Plans",
            "entity_id": plan_id,
        },
    )

    assert status == 200
    assert _payload(body)["status"] == "REJECTED"
    assert (plan_folder / "approval.pdf").read_bytes() == b"keep evidence"


def test_evidence_preview_post_returns_pdf_and_image_bytes_as_json(evidence_host):
    base_url, server = evidence_host
    record_dir = server.EVIDENCE_ROOT / "TargetPlan" / "PLN-PREVIEW-01"
    record_dir.mkdir(parents=True)
    cases = [
        ("agreement.pdf", b"%PDF-1.7 preview payload", "application/pdf"),
        ("site-photo.jpg", b"jpeg preview payload", "image/jpeg"),
    ]

    for filename, original_bytes, mime_type in cases:
        (record_dir / filename).write_bytes(original_bytes)
        status, headers, body = _post_json(
            base_url,
            "/evidence/preview",
            {
                "module": "TargetPlan",
                "record_id": "PLN-PREVIEW-01",
                "filename": filename,
            },
        )

        assert status == 200
        assert headers.get_content_type() == "application/json"
        assert "Content-Disposition" not in headers
        payload = _payload(body)
        assert payload["status"] == "READY"
        assert payload["mime_type"] == mime_type
        assert payload["filename"] == filename
        assert base64.b64decode(payload["data"], validate=True) == original_bytes


@pytest.mark.parametrize(
    ("payload", "expected_status"),
    [
        ({"module": "Plans", "record_id": "PLN-01", "filename": "file.pdf"}, 400),
        ({"module": "TargetPlan", "record_id": "../outside", "filename": "file.pdf"}, 400),
        ({"module": "TargetPlan", "record_id": "PLN-01", "filename": "../outside.pdf"}, 400),
        ({"module": "TargetPlan", "record_id": "PLN-01", "filename": "file.docx"}, 400),
        ({"module": "TargetPlan", "record_id": "PLN-01", "filename": "missing.pdf"}, 404),
    ],
)
def test_evidence_preview_post_rejects_invalid_or_missing_files(evidence_host, payload, expected_status):
    base_url, _server = evidence_host

    status, headers, body = _post_json(base_url, "/evidence/preview", payload)

    assert status == expected_status
    assert headers.get_content_type() == "application/json"
    result = _payload(body)
    assert result["status"] == "REJECTED"
    assert result["errors"]


def test_evidence_listing_returns_only_supported_direct_child_files(evidence_host):
    base_url, server = evidence_host
    record_dir = server.EVIDENCE_ROOT / "TargetPlan" / "PLN-TEST-01"
    record_dir.mkdir(parents=True)
    (record_dir / "agreement.pdf").write_bytes(b"%PDF-1.4 test")
    (record_dir / "field-photo.JPG").write_bytes(b"jpeg test")
    (record_dir / "approval.docx").write_bytes(b"word test")
    (record_dir / "not-evidence.exe").write_bytes(b"ignore")

    status, headers, body = _get(
        base_url,
        "/evidence/list",
        {"module": "TargetPlan", "record_id": "PLN-TEST-01"},
    )

    assert status == 200
    assert headers.get_content_type() == "application/json"
    payload = _payload(body)
    assert payload["status"] == "READY"
    assert payload["module"] == "TargetPlan"
    assert payload["record_id"] == "PLN-TEST-01"
    assert payload["folder"] == "Database/Evidence/TargetPlan/PLN-TEST-01"

    files = {item["name"]: item for item in payload["files"]}
    assert set(files) == {"agreement.pdf", "field-photo.JPG", "approval.docx"}
    assert files["agreement.pdf"]["previewable"] is True
    assert files["field-photo.JPG"]["previewable"] is True
    assert files["approval.docx"]["previewable"] is False
    assert all(item["available"] is True for item in files.values())


@pytest.mark.parametrize(
    ("filename", "body", "expected_type", "expected_disposition"),
    [
        ("agreement.pdf", b"%PDF-1.4 test", "application/pdf", "inline"),
        ("field-photo.jpg", b"jpeg test", "image/jpeg", "inline"),
        ("approval.docx", b"word test", None, "attachment"),
    ],
)
def test_evidence_file_endpoint_previews_images_and_downloads_word_documents(
    evidence_host, filename, body, expected_type, expected_disposition
):
    base_url, server = evidence_host
    record_dir = server.EVIDENCE_ROOT / "TargetPlan" / "PLN-TEST-02"
    record_dir.mkdir(parents=True)
    (record_dir / filename).write_bytes(body)

    status, headers, response_body = _get(
        base_url,
        "/evidence/file",
        {
            "module": "TargetPlan",
            "record_id": "PLN-TEST-02",
            "filename": filename,
        },
    )

    assert status == 200
    assert response_body == body
    if expected_type:
        assert headers.get_content_type() == expected_type
    assert headers.get("Content-Disposition", "").startswith(expected_disposition + ";")
    assert headers.get("Cache-Control") == "no-store"
    assert headers.get("X-Content-Type-Options") == "nosniff"


@pytest.mark.parametrize(
    ("endpoint", "params"),
    [
        ("/evidence/list", {"module": "Plans", "record_id": "PLN-01"}),
        ("/evidence/list", {"module": "TargetPlan", "record_id": "../outside"}),
        (
            "/evidence/file",
            {"module": "TargetPlan", "record_id": "PLN-01", "filename": "../outside.pdf"},
        ),
        (
            "/evidence/file",
            {"module": "TargetPlan", "record_id": "PLN-01", "filename": r"..\\outside.pdf"},
        ),
        (
            "/evidence/file",
            {"module": "TargetPlan", "record_id": "PLN-01", "filename": "payload.exe"},
        ),
    ],
)
def test_evidence_endpoint_rejects_invalid_module_ids_paths_and_extensions(
    evidence_host, endpoint, params
):
    base_url, _server = evidence_host

    status, _headers, body = _get(base_url, endpoint, params)

    assert status == 400
    payload = _payload(body)
    assert payload["status"] == "REJECTED"
    assert payload["errors"]


def test_evidence_file_endpoint_returns_not_found_for_missing_file(evidence_host):
    base_url, _server = evidence_host

    status, _headers, body = _get(
        base_url,
        "/evidence/file",
        {
            "module": "TargetPlan",
            "record_id": "PLN-MISSING-01",
            "filename": "missing.pdf",
        },
    )

    assert status == 404
    payload = _payload(body)
    assert payload["status"] == "REJECTED"
    assert payload["errors"][0]["code"] == "HOST-001"
