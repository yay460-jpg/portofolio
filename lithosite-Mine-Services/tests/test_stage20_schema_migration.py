import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))

from mine_services.schema import HEADERS as A3_HEADERS
from mine_services.schema_migration import SchemaMigrationError, migrate_a2_to_a3
from mine_services.schema_migration_contract import A2_HEADERS, A2_SHEETS, A3_SHEETS


def make_a2_workbook(path: Path):
    wb = Workbook()
    first = wb.active
    wb.remove(first)

    for sheet in A2_SHEETS:
        ws = wb.create_sheet(sheet)
        if sheet in A2_HEADERS:
            ws.append(A2_HEADERS[sheet])
        elif sheet == "_System":
            ws.append(["schema_version", "A.2"])
        elif sheet == "_Lists":
            ws.append(["equipment_category"])
            ws.append(["Heavy Equipment"])
        elif sheet == "_Baseline":
            ws.append(["baseline"])
            ws.append(["A.2 baseline"])
    wb["Equipment"].append([
        "EQ-MIG-001", "DT-MIG-001", "Heavy Equipment", "Dump Truck",
        "Owner", None, "Active", None, None
    ])
    wb["HSE"].append([
        "HSE-MIG-001", None, "Road & Hauling", None, "Incident",
        "High", "migration test", None, "Open", None
    ])
    wb["AuditLog"].append([
        "AUD-MIG-001", None, "HSE", "HSE-MIG-001", "CREATE",
        None, "preserve", "test"
    ])
    wb.save(path)


def test_stage20_migration_creates_a3_without_mutating_a2(tmp_path):
    source = tmp_path / "source-a2.xlsx"
    target = tmp_path / "target-a3.xlsx"
    make_a2_workbook(source)

    before = load_workbook(source, data_only=False)
    before_sheets = list(before.sheetnames)
    before_equipment = list(before["Equipment"].values)
    before_hse = list(before["HSE"].values)

    result = migrate_a2_to_a3(source, target)

    assert result == target
    assert target.exists()

    after_source = load_workbook(source, data_only=False)
    assert list(after_source.sheetnames) == before_sheets
    assert list(after_source["Equipment"].values) == before_equipment
    assert list(after_source["HSE"].values) == before_hse

    migrated = load_workbook(target, data_only=False)
    assert list(migrated.sheetnames) == A3_SHEETS
    assert migrated["_System"]["B1"].value == "A.3"
    assert list(migrated["MapMarker"].values) == [(
        "marker_id", "marker_type", "label", "easting", "northing",
        "elevation", "source_entity", "source_id", "status"
    )]
    assert migrated["MapMarker"].max_row == 1
    assert list(migrated["GlobalCapacity"].values)[0] == tuple(A3_HEADERS["GlobalCapacity"])
    assert list(migrated["Checker"].values)[0] == tuple(A3_HEADERS["Checker"])
    assert list(migrated["WorkFront"].values)[0] == tuple(A3_HEADERS["WorkFront"])
    assert list(migrated["Operations"].values)[0] == tuple(A3_HEADERS["Operations"])
    migrated_plan_headers = list(migrated["Plans"].values)[0]
    expected_plan_headers = tuple(
        "measurement" if header == "unit" else header
        for header in A2_HEADERS["Plans"]
        if header != "target_hours"
    )
    assert migrated_plan_headers == expected_plan_headers
    assert "target_hours" not in migrated_plan_headers


def test_stage20_migration_preserves_existing_domain_data_and_audit(tmp_path):
    source = tmp_path / "source-a2.xlsx"
    target = tmp_path / "target-a3.xlsx"
    make_a2_workbook(source)

    migrate_a2_to_a3(source, target)
    wb = load_workbook(target, data_only=True)

    assert list(wb["Equipment"].values)[1][0] == "EQ-MIG-001"
    assert list(wb["HSE"].values)[1][0] == "HSE-MIG-001"
    assert list(wb["AuditLog"].values)[1][0] == "AUD-MIG-001"


def test_stage20_migration_rejects_non_a2_source(tmp_path):
    source = tmp_path / "wrong-version.xlsx"
    target = tmp_path / "target-a3.xlsx"
    make_a2_workbook(source)

    wb = load_workbook(source)
    wb["_System"]["B1"] = "A.1"
    wb.save(source)

    with pytest.raises(SchemaMigrationError, match="SCHEMA_VERSION"):
        migrate_a2_to_a3(source, target)

    assert not target.exists()


def test_stage20_migration_rejects_existing_target(tmp_path):
    source = tmp_path / "source-a2.xlsx"
    target = tmp_path / "target-a3.xlsx"
    make_a2_workbook(source)
    target.write_bytes(b"existing")

    with pytest.raises(FileExistsError):
        migrate_a2_to_a3(source, target)

    assert target.read_bytes() == b"existing"


def test_stage20_migration_rejects_in_place_source_target(tmp_path):
    source = tmp_path / "source-a2.xlsx"
    make_a2_workbook(source)

    with pytest.raises(SchemaMigrationError, match="SOURCE_TARGET_MUST_DIFFER"):
        migrate_a2_to_a3(source, source)
